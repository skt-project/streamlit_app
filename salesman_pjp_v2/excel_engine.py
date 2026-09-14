"""
Excel template generation — direct port of the production logic.
No business logic changes; only imports updated to use local config.
"""

import re
import unicodedata
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .config import PJP_COLS, DAY_OPTIONS, WEEK_OPTIONS, FREQUENCY_OPTIONS


# ─── Style helpers ────────────────────────────────────────────────────────────

def _thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _header_font():
    return Font(bold=True, color="FFFFFF", size=10, name="Calibri")

def _note_font():
    return Font(italic=True, color="808080", size=9, name="Calibri")

def _req_font():
    return Font(bold=True, color="C00000", size=9, name="Calibri")

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _vcenter(wrap=False):
    return Alignment(vertical="center", wrap_text=wrap)


# ─── Named-range key sanitiser ────────────────────────────────────────────────

def _safe_name(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", text)
    ascii_s = nfkd.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9]", "_", ascii_s)
    if not cleaned or cleaned[0].isdigit():
        cleaned = "_" + cleaned
    return ("NR_" + cleaned)[:255]


def _indirect_clean(cell_ref: str) -> str:
    special = [" ", "-", "/", "(", ")", "+", "&", ".", "'"]
    expr = cell_ref
    for ch in special:
        expr = f'SUBSTITUTE({expr},"{ch}","_")'
    return expr


# ─── Lookup sheet + named ranges ──────────────────────────────────────────────

def _build_lookup_and_named_ranges(wb: Workbook, dist_df: pd.DataFrame,
                                   store_df: pd.DataFrame | None = None) -> None:
    LK = "Lookup"
    lk = wb.create_sheet(LK)
    lk.sheet_state = "hidden"

    asm_list = sorted(dist_df["asm"].dropna().unique().tolist())
    cur_col = 1

    # ASM list
    lk.cell(row=1, column=cur_col, value="__ALL_ASM__")
    for i, asm in enumerate(asm_list, start=2):
        lk.cell(row=i, column=cur_col, value=asm)
    c = get_column_letter(cur_col)
    nm = _safe_name("ALL_ASM")
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{LK}'!${c}$2:${c}${1+len(asm_list)}")
    cur_col += 1

    # Per-ASM region lists
    for asm in asm_list:
        regions = sorted(dist_df.loc[dist_df["asm"] == asm, "region"].unique().tolist())
        lk.cell(row=1, column=cur_col, value=f"__ASM_{asm}__")
        for i, reg in enumerate(regions, start=2):
            lk.cell(row=i, column=cur_col, value=reg)
        c = get_column_letter(cur_col)
        nm = _safe_name(asm)
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{LK}'!${c}$2:${c}${1+len(regions)}")
        cur_col += 1

    # Per-ASM+Region distributor lists
    for asm in asm_list:
        regions = sorted(dist_df.loc[dist_df["asm"] == asm, "region"].unique().tolist())
        for region in regions:
            mask = (dist_df["asm"] == asm) & (dist_df["region"] == region)
            names = sorted(dist_df.loc[mask, "distributor_name"].unique().tolist())
            lk.cell(row=1, column=cur_col, value=f"__ASM_{asm}__REG_{region}__")
            for i, name in enumerate(names, start=2):
                lk.cell(row=i, column=cur_col, value=name)
            c = get_column_letter(cur_col)
            nm = _safe_name(f"{asm}_{region}")
            wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{LK}'!${c}$2:${c}${1+len(names)}")
            cur_col += 1

    # Distributor name → code lookup
    name_col, code_col = cur_col, cur_col + 1
    lk.cell(row=1, column=name_col, value="__DIST_NAME__")
    lk.cell(row=1, column=code_col, value="__DIST_CODE__")
    all_dists = (
        dist_df[["distributor_name", "distributor_code"]]
        .drop_duplicates(subset=["distributor_name"])
        .sort_values("distributor_name")
        .reset_index(drop=True)
    )
    for i, row in all_dists.iterrows():
        lk.cell(row=i + 2, column=name_col, value=row["distributor_name"])
        lk.cell(row=i + 2, column=code_col, value=row["distributor_code"])
    nc, kc = get_column_letter(name_col), get_column_letter(code_col)
    last_row = 1 + len(all_dists)
    wb.defined_names["NR_DIST_LOOKUP"] = DefinedName(
        "NR_DIST_LOOKUP", attr_text=f"'{LK}'!${nc}$2:${kc}${last_row}"
    )
    cur_col += 2

    # Per-distributor store lists
    if store_df is not None and not store_df.empty:
        for dist_name in sorted(store_df["distributor_name"].dropna().unique().tolist()):
            codes = sorted(
                store_df.loc[store_df["distributor_name"] == dist_name, "store_code"]
                .dropna().unique().tolist()
            )
            lk.cell(row=1, column=cur_col, value=f"__STORE_{dist_name}__")
            for i, code in enumerate(codes, start=2):
                lk.cell(row=i, column=cur_col, value=code)
            c = get_column_letter(cur_col)
            nm = _safe_name(f"STORE_{dist_name}")
            wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{LK}'!${c}$2:${c}${1+len(codes)}")
            cur_col += 1

        sc_col, sn_col = cur_col, cur_col + 1
        lk.cell(row=1, column=sc_col, value="__STORE_CODE__")
        lk.cell(row=1, column=sn_col, value="__STORE_NAME__")
        all_stores = (
            store_df[["store_code", "store_name"]]
            .drop_duplicates(subset=["store_code"])
            .sort_values("store_code")
            .reset_index(drop=True)
        )
        for i, row in all_stores.iterrows():
            lk.cell(row=i + 2, column=sc_col, value=row["store_code"])
            lk.cell(row=i + 2, column=sn_col, value=row["store_name"])
        scc, snc = get_column_letter(sc_col), get_column_letter(sn_col)
        last_row = 1 + len(all_stores)
        wb.defined_names["NR_STORE_LOOKUP"] = DefinedName(
            "NR_STORE_LOOKUP", attr_text=f"'{LK}'!${scc}$2:${snc}${last_row}"
        )


def _attach_cascade_dvs(ws, col_names: list, first_data: int, last_data: int) -> None:
    def cl(name):
        return get_column_letter(col_names.index(name) + 1)

    def sqref(name):
        c = cl(name)
        return f"{c}{first_data}:{c}{last_data}"

    asm_ref = f"{cl('ASM')}{first_data}"
    reg_ref = f"{cl('Region')}{first_data}"

    dv_asm = DataValidation(type="list", formula1=_safe_name("ALL_ASM"), allow_blank=True,
        showInputMessage=True, promptTitle="Langkah 1 - ASM",
        prompt="Pilih nama ASM. Region dan Distributor akan menyesuaikan.",
        showErrorMessage=True, errorTitle="Input Tidak Valid", error="Pilih ASM dari daftar.")
    ws.add_data_validation(dv_asm)
    dv_asm.sqref = sqref("ASM")

    asm_clean = _indirect_clean(asm_ref)
    dv_reg = DataValidation(type="list", formula1=f'INDIRECT("NR_"&{asm_clean})',
        allow_blank=True, showInputMessage=True, promptTitle="Langkah 2 - Region",
        prompt="Pilih Region. Daftar disesuaikan dengan ASM.",
        showErrorMessage=True, errorTitle="Input Tidak Valid",
        error="Pilih Region dari daftar. Pastikan ASM sudah dipilih.")
    ws.add_data_validation(dv_reg)
    dv_reg.sqref = sqref("Region")

    reg_clean = _indirect_clean(reg_ref)
    dv_nama = DataValidation(type="list",
        formula1=f'INDIRECT("NR_"&{asm_clean}&"_"&{reg_clean})',
        allow_blank=True, showInputMessage=True, promptTitle="Langkah 3 - Nama Distributor",
        prompt="Pilih Nama Distributor. Daftar disesuaikan dengan ASM dan Region.",
        showErrorMessage=True, errorTitle="Input Tidak Valid",
        error="Pilih Distributor dari daftar. Pastikan ASM dan Region sudah dipilih.")
    ws.add_data_validation(dv_nama)
    dv_nama.sqref = sqref("Nama Distributor")

    if "Kode Toko" in col_names:
        dist_ref = f"{cl('Nama Distributor')}{first_data}"
        dist_clean = _indirect_clean(dist_ref)
        dv_store = DataValidation(type="list",
            formula1=f'INDIRECT("NR_STORE_"&{dist_clean})',
            allow_blank=True, showInputMessage=True, promptTitle="Langkah 4 - Kode Toko",
            prompt="Pilih Kode Toko. Daftar disesuaikan dengan Distributor.",
            showErrorMessage=True, errorTitle="Input Tidak Valid",
            error="Pilih Kode Toko dari daftar. Pastikan Nama Distributor sudah dipilih.")
        ws.add_data_validation(dv_store)
        dv_store.sqref = sqref("Kode Toko")


# ─── Main Excel generator ─────────────────────────────────────────────────────

def create_pjp_excel(df: pd.DataFrame, distributor_map: dict,
                     dist_df: pd.DataFrame, store_df: pd.DataFrame) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _build_lookup_and_named_ranges(wb, dist_df, store_df)

    col_names  = [c for c, _, _ in PJP_COLS]
    col_types  = {c: t for c, _, t in PJP_COLS}
    col_req    = {c: r for c, r, _ in PJP_COLS}
    CASCADE    = {"ASM", "Region", "Nama Distributor", "Kode Distributor", "Kode Toko", "Nama Toko"}
    FIRST_DATA = 4
    LAST_DATA  = 30003

    notes = {
        "ASM":                "Langkah 1 - Pilih ASM dari dropdown",
        "Region":             "Langkah 2 - Pilih Region (mengikuti ASM)",
        "Nama Distributor":   "Langkah 3 - Pilih Distributor (mengikuti Region)",
        "Kode Distributor":   "Otomatis terisi dari Nama Distributor",
        "Nama Salesman":      "Teks bebas",
        "Kode Toko":          "Langkah 4 - Pilih Kode Toko (mengikuti Distributor)",
        "Nama Toko":          "Otomatis terisi dari Kode Toko",
        "Hari":               "Drop down dengan opsi hari",
        "Minggu Ganjil/Minggu Genap/Minggu Ganjil + Genap":
                              "Drop down: ganjil / genap / ganjil+genap",
        "Frekuensi":          "F4+ = >1x seminggu  |  F4 = 1x/minggu  |  F2 = 2x/bulan  |  F1 = 1x/bulan",
    }

    ws = wb.create_sheet("PJP Template")

    for ci, cn in enumerate(col_names, 1):
        cell = ws.cell(row=1, column=ci, value=notes.get(cn, ""))
        cell.font = _note_font()
        cell.alignment = _vcenter(wrap=True)

    for ci, cn in enumerate(col_names, 1):
        if col_req.get(cn):
            cell = ws.cell(row=2, column=ci, value="Wajib Diisi")
            cell.font = _req_font()
            cell.alignment = _center()

    for ci, cn in enumerate(col_names, 1):
        cell = ws.cell(row=3, column=ci, value=cn)
        cell.font = _header_font()
        cell.fill = _fill("1A7A6E" if cn in CASCADE else "ED7D31")
        cell.alignment = _center()
        cell.border = _thin_border()

    ws.row_dimensions[1].height = 42
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 44
    ws.freeze_panes = "A4"
    for ci, w in enumerate([22, 24, 30, 20, 22, 18, 30, 12, 40, 22], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    def col_letter(name):
        return get_column_letter(col_names.index(name) + 1)

    def dr(name):
        c = col_letter(name)
        return f"{c}{FIRST_DATA}:{c}{LAST_DATA}"

    _attach_cascade_dvs(ws, col_names, FIRST_DATA, LAST_DATA)

    for col_name, opts in [
        ("Hari", DAY_OPTIONS),
        ("Minggu Ganjil/Minggu Genap/Minggu Ganjil + Genap", WEEK_OPTIONS),
    ]:
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"',
            allow_blank=True, showInputMessage=True, promptTitle=col_name,
            prompt=f"Pilih {col_name}", showErrorMessage=True,
            errorTitle="Input Tidak Valid", error="Pilih nilai dari daftar dropdown.")
        ws.add_data_validation(dv)
        dv.sqref = dr(col_name)

    freq_dv = DataValidation(type="list",
        formula1='"' + ",".join(FREQUENCY_OPTIONS) + '"',
        allow_blank=True, showInputMessage=True, promptTitle="Frekuensi Kunjungan",
        prompt="F4+ = >1x/minggu | F4 = 1x/minggu | F2 = 2x/bulan | F1 = 1x/bulan",
        showErrorMessage=True, errorTitle="Input Tidak Valid",
        error="Pilih F4+, F4, F2, atau F1.")
    ws.add_data_validation(freq_dv)
    freq_dv.sqref = dr("Frekuensi")

    nama_cl     = col_letter("Nama Distributor")
    kode_toko_cl = col_letter("Kode Toko")
    df_reindexed = df.reindex(columns=col_names)

    for excel_row in range(FIRST_DATA, LAST_DATA + 1):
        dfi = excel_row - FIRST_DATA
        has_data = dfi < len(df_reindexed)

        for ci, cn in enumerate(col_names, 1):
            cell = ws.cell(row=excel_row, column=ci)

            if cn == "Kode Distributor":
                cell.value = f'=IFERROR(VLOOKUP({nama_cl}{excel_row},NR_DIST_LOOKUP,2,0),"")'
                cell.fill = _fill("D6E4F0")
                cell.font = Font(italic=True, color="1A7A6E", size=10, name="Calibri")
                cell.alignment = _vcenter()
                cell.border = _thin_border()
                cell.number_format = "@"
                cell.protection = Protection(locked=True)
                continue

            if cn == "Nama Toko":
                cell.value = f'=IFERROR(VLOOKUP({kode_toko_cl}{excel_row},NR_STORE_LOOKUP,2,0),"")'
                cell.fill = _fill("D6E4F0")
                cell.font = Font(italic=True, color="1A7A6E", size=10, name="Calibri")
                cell.alignment = _vcenter()
                cell.border = _thin_border()
                cell.number_format = "@"
                cell.protection = Protection(locked=True)
                continue

            cell.number_format = "@"
            if has_data:
                val = df_reindexed.iloc[dfi].get(cn, "")
                cell.value = "" if pd.isna(val) else (str(val) if val != "" else "")
            cell.alignment = _vcenter()
            cell.border = _thin_border()
            cell.protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.protection.password = "skintific"
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
