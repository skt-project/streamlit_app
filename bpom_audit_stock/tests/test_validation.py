"""Tes validasi Excel: python -m pytest tests"""
import sys
from io import BytesIO
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import validation as v  # noqa: E402

NAME = "CV ANDIRA DAHAYU - BENGKULU"
TEMPLATE = Path(__file__).resolve().parent.parent / "template" / "Contoh Excel Response.xlsx"


def make(rows, headers=v.EXPECTED_HEADERS, title=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    if title:
        ws.append([None, None, None, None, "Filled by Distributor"])
    ws.append(list(headers))
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def good(n=1, qty=5, pid="P-1", bpom="NA1"):
    return [NAME, pid, "Produk", bpom, qty]


def messages(res):
    return " | ".join(i.message for i in res.issues)


def test_valid_file_with_title_row():
    res = v.parse_excel(make([good(), good(qty=0, bpom="NA2"), good(qty=-3, bpom="NA3")]), NAME)
    assert res.ok and [r.qty for r in res.rows] == [5, 0, -3]
    assert res.rows[0].excel_row == 3


def test_name_match_ignores_case_and_spacing():
    res = v.parse_excel(make([["cv andira dahayu-bengkulu", "P", "N", "B", 1]]), NAME)
    assert res.ok


def test_header_case_sensitive():
    h = list(v.EXPECTED_HEADERS); h[4] = "Qty"
    res = v.parse_excel(make([good()], headers=h), NAME)
    assert not res.ok and "persis" in messages(res)


def test_extra_missing_and_reordered_columns():
    assert not v.parse_excel(make([good() + ["x"]], headers=[*v.EXPECTED_HEADERS, "EXTRA"]), NAME).ok
    assert not v.parse_excel(make([good()[:4]], headers=v.EXPECTED_HEADERS[:4]), NAME).ok
    swapped = [v.EXPECTED_HEADERS[1], v.EXPECTED_HEADERS[0], *v.EXPECTED_HEADERS[2:]]
    assert "Urutan" in messages(v.parse_excel(make([good()], headers=swapped), NAME))


def test_qty_rules():
    for bad in (None, "", 1.5, "abc", "1,5", True):
        assert v.parse_qty(bad)[0] is None, bad
    assert v.parse_qty(7.0) == (7, None)
    assert v.parse_qty("-12") == (-12, None)
    assert v.parse_qty(0) == (0, None)


def test_all_or_nothing_and_row_numbers():
    res = v.parse_excel(make([good(), good(qty=None, bpom="NA2"), good(bpom="NA3")]), NAME)
    assert not res.ok
    assert [i.row for i in res.issues] == [4]


def test_duplicate_key_within_file():
    res = v.parse_excel(make([good(), good(pid="p-1", bpom="na1")]), NAME)
    assert not res.ok and "Duplikat dengan baris 3" in messages(res)


def test_required_text_fields_and_wrong_distributor():
    res = v.parse_excel(make([["CV LAIN - X", "", "N", "B", 1]]), NAME)
    cols = {i.column for i in res.issues}
    assert v.EXPECTED_HEADERS[0] in cols and "PRODUCT ID" in cols


def test_blank_rows_skipped_and_empty_file_rejected():
    assert v.parse_excel(make([good(), [None] * 5, [None] * 5]), NAME).ok
    assert not v.parse_excel(make([]), NAME).ok


def test_not_xlsx():
    assert not v.parse_excel(b"bukan excel", NAME).ok


SKUS = [v.Sku("P-1", "Produk", "NA1"), v.Sku("P-1", "Produk", "NA2"), v.Sku("P-2", "Lain", "NB1")]


def test_sku_rules_known_additional_and_missing():
    rows = [good(), good(bpom="NA2"), [NAME, "P-1", "Produk", "NA-BARU", 3]]
    res = v.parse_excel(make(rows), NAME, SKUS)
    assert res.ok and [r.additional for r in res.rows] == [False, False, True]
    assert [k.product_id for k in res.missing] == ["P-2"]


def test_unknown_sku_and_name_mismatch_rejected():
    res = v.parse_excel(make([[NAME, "P-9", "X", "NA1", 1]]), NAME, SKUS)
    assert not res.ok and any(i.column == "PRODUCT ID" for i in res.issues)
    res = v.parse_excel(make([[NAME, "P-1", "Nama Lain", "NA1", 1]]), NAME, SKUS)
    assert not res.ok and any(i.column == "PRODUCT NAME" for i in res.issues)


def test_form_validation_has_no_period():
    assert v.validate_form(company="A", distributor="B", tanggal="d", pic="x", has_file=True) == []
    assert len(v.validate_form(company=None, distributor=None, tanggal=None, pic="", has_file=False)) == 5


def test_reference_files_load():
    import reference_data as r
    ref = Path(__file__).resolve().parent.parent / "reference"
    dist = r.load_distributors(ref / "Daftar Distributor.xlsx")
    skus = r.load_skus(ref / "Data SKU BPOM.xlsx")
    assert len(dist) == 94 and len(skus) == 71


def test_template_is_built_from_sku_list_and_original_untouched():
    import reference_data as r
    from template_builder import build_template
    ref = Path(__file__).resolve().parent.parent / "reference"
    skus = r.load_skus(ref / "Data SKU BPOM.xlsx")
    raw = TEMPLATE.read_bytes()
    out = build_template(raw, NAME, skus)
    assert raw == TEMPLATE.read_bytes()
    ws = openpyxl.load_workbook(BytesIO(out)).active
    assert [c.value for c in ws[2]] == list(v.EXPECTED_HEADERS) and ws["E1"].value == "Filled by Distributor"
    assert ws.max_row == 2 + len(skus)
    for i, sku in enumerate(skus, start=3):
        assert [c.value for c in ws[i]] == [NAME, sku.product_id, sku.product_name, sku.bpom_number, None]
    for i in range(3, ws.max_row + 1):   # isi QTY -> lolos validasi terhadap daftar SKU
        ws.cell(i, 5).value = i
    buf = BytesIO(); ws.parent.save(buf)
    res = v.parse_excel(buf.getvalue(), NAME, skus)
    assert res.ok and len(res.rows) == len(skus) and not res.missing
