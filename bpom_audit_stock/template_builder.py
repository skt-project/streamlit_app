"""Membuat template untuk satu distributor dari template resmi + daftar SKU/BPOM prinsipal.

File template asli tidak pernah diubah: file dibaca ke memori, baris data diganti sesuai daftar
SKU terkini dan kolom A diisi nama distributor, lalu disimpan sebagai bytes baru. Header, baris judul,
lebar kolom, dan gaya sel tetap seperti template.
"""
from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Sequence

import openpyxl

from validation import EXPECTED_HEADERS, HEADER_SEARCH_ROWS, Sku, _looks_like_header


def build_template(template_bytes: bytes, distributor_name: str, skus: Sequence[Sku]) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(template_bytes))
    ws = wb.worksheets[0]
    header_row = next(
        (no for no, cells in enumerate(ws.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True), start=1)
         if _looks_like_header(cells)),
        None,
    )
    if header_row is None:
        raise ValueError("Header template tidak ditemukan.")

    ncols = len(EXPECTED_HEADERS)
    first = header_row + 1
    styles = [copy(ws.cell(first, c)._style) for c in range(1, ncols + 1)]   # gaya baris data pertama
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)

    for i, sku in enumerate(skus):
        values = (distributor_name, sku.product_id, sku.product_name, sku.bpom_number, None)
        for c, value in enumerate(values, start=1):
            cell = ws.cell(first + i, c, value)
            cell._style = copy(styles[c - 1])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
