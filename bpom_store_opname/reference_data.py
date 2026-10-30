"""Data referensi dari BD: daftar toko (`reference/Daftar Store.xlsx`) serta katalog SKU + nomor BPOM (Sheet1 template resmi)."""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from validation import (EXPECTED_HEADERS, HEADER_SEARCH_ROWS, LEGACY_FIRST_HEADER, Sku, normalize_name, row_key)

# Kolom daftar toko dari BD. Nama kolom dibandingkan tanpa peduli huruf besar/kecil dan spasi; kolom lain diabaikan.
REGION_COLS = ("region",)
GROUP_COLS = ("nama group toko", "nama grup toko", "group toko", "grup toko", "nama group", "nama grup", "group", "grup")
TOKO_COLS = ("nama toko", "toko", "store name")
STORE_HEADERS_HELP = ("Region", "Nama Group Toko", "Nama Toko")      # untuk pesan error


class ReferenceDataError(Exception):
    """File referensi rusak / tidak sesuai format."""


class StoreListMissing(ReferenceDataError):
    """Daftar toko dari BD belum diletakkan di reference/ (bukan file rusak)."""


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _col_key(value) -> str:
    return re.sub(r"\s+", " ", _clean(value)).casefold()


@dataclass(frozen=True)
class Toko:
    """Satu pilihan di dropdown: nama grup bila toko milik grup, selain itu nama toko ("Nama Toko (By Group)")."""
    region: str
    name: str                  # nama grup bila ada, selain itu nama toko; ini yang masuk kolom A template dan ke sheet
    is_group: bool = False
    outlets: int = 1           # jumlah baris toko di daftar BD yang digabung ke pilihan ini
    search_text: str = ""      # huruf kecil, untuk kotak cari (dihitung sekali saat memuat)

    @property
    def label(self) -> str:
        return f"{self.name}  ·  grup ({self.outlets} toko)" if self.is_group and self.outlets > 1 else self.name


@dataclass(frozen=True)
class RowFilter:
    """Hanya baris yang nilai kolomnya SAMA PERSIS (tanpa peduli huruf besar/kecil dan spasi) dengan `equals` yang dipakai.

    Kolom dicari berdasarkan nama header (`header_names`); bila tidak ada header itu, dipakai huruf kolom `column_letter`
    (mis. "H"). Baris yang tidak lolos filter dilewati sepenuhnya, termasuk validasinya.
    """
    equals: str
    header_names: tuple[str, ...] = ()
    column_letter: str = ""

    def resolve(self, header: list[str]) -> int:
        for name in self.header_names:
            if _col_key(name) in header:
                return header.index(_col_key(name))
        if self.column_letter:
            index = _letter_index(self.column_letter)
            if index < len(header):
                return index
        raise ReferenceDataError(
            f"Kolom filter tidak ditemukan (header {list(self.header_names)} atau kolom {self.column_letter or '-'}); "
            f"header yang ada: {header}.")


def _index_letter(index: int) -> str:
    """Indeks kolom berbasis 0 -> huruf kolom Excel/Sheets (0 -> A, 7 -> H, 26 -> AA)."""
    n, letters = index + 1, ""
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters


def _letter_index(letter: str) -> int:
    n = 0
    for ch in letter.strip().upper():
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def tokos_from_rows(rows, source: str, *, row_filter: RowFilter | None = None, max_listed: int = 8) -> list[Toko]:
    """Susun pilihan toko dari baris-baris tabel (iterator tuple/list nilai sel). `source` hanya untuk pesan error.

    Kolom wajib: Region dan Nama Toko; Nama Group Toko opsional (urutan bebas, kolom lain diabaikan). Baris dengan grup terisi
    menjadi SATU pilihan atas nama grup; baris tanpa grup menjadi pilihan atas nama toko; unik per (Region, nama).
    Satu masalah saja menggagalkan seluruh pemuatan supaya tidak ada toko yang diam-diam hilang.
    Header boleh di baris mana pun dalam HEADER_SEARCH_ROWS baris pertama.
    """
    problems: list[str] = []
    merged: dict[tuple[str, str], list] = {}      # (region, nama) -> [Region asli, nama asli, is_group, jumlah baris]
    rows = iter(rows)
    header, header_no, cols = [], 0, None
    for no, cells in enumerate(rows, start=1):
        candidate = [_col_key(c) for c in (cells or ())]
        if not header:
            header, header_no = candidate, no
        found = _find_store_columns(candidate)
        if found:
            header, header_no, cols = candidate, no, found
            break
        if no >= HEADER_SEARCH_ROWS:
            break
    if cols is None:
        raise ReferenceDataError(
            f"Header '{source}' tidak ditemukan. Perlu kolom Region dan Nama Toko (opsional Nama Group Toko); "
            f"ditemukan {header}.")
    i_region, i_group, i_toko = cols
    i_filter = row_filter.resolve(header) if row_filter else None
    wanted = _col_key(row_filter.equals) if row_filter else ""
    seen_values: Counter = Counter()           # nilai yang ada di kolom filter (untuk pesan error bila tidak ada yang cocok)
    for no, row in enumerate(rows, start=header_no + 1):
        row = tuple(row) + (None,) * (len(header) - len(row))
        if all(_clean(c) == "" for c in row):
            continue
        if row_filter:
            seen_values[_clean(row[i_filter])] += 1
            if _col_key(row[i_filter]) != wanted:
                continue
        region, toko = _clean(row[i_region]), _clean(row[i_toko])
        group = _clean(row[i_group]) if i_group is not None else ""
        name = group or toko
        if not (region and name):
            problems.append(f"baris {no}: Region dan Nama Toko wajib diisi")
            continue
        key = (normalize_name(region), normalize_name(name))
        entry = merged.setdefault(key, [region, name, bool(group), 0])
        entry[3] += 1
        if bool(group) != entry[2]:
            problems.append(f"baris {no}: '{name}' di region '{region}' muncul sebagai grup dan sebagai toko tunggal")
    if problems:
        shown = "; ".join(problems[:max_listed]) + (
            f"; ... (+{len(problems) - max_listed} masalah lain)" if len(problems) > max_listed else "")
        raise ReferenceDataError(f"'{source}' bermasalah: {shown}.")
    if not merged:
        if row_filter:
            col = f"'{header[i_filter]}' (kolom {_index_letter(i_filter)})"
            values = ", ".join(f"{v!r} ({n})" for v, n in seen_values.most_common(8)) or "(kolom kosong)"
            raise ReferenceDataError(
                f"'{source}' tidak berisi toko dengan {_clean(row_filter.equals)} pada kolom filter {col}. "
                f"Nilai yang ada di kolom itu: {values}. Bila kolomnya salah, atur store_filter_header / store_filter_column; "
                "bila nilainya berbeda, atur store_filter_value.")
        raise ReferenceDataError(f"'{source}' tidak berisi toko.")
    return [Toko(region, name, is_group, n, name.lower()) for region, name, is_group, n in merged.values()]


def load_stores(path: Path, *, row_filter: RowFilter | None = None, max_listed: int = 8) -> list[Toko]:
    """Daftar toko dari file xlsx BD (`reference/Daftar Store.xlsx`); aturan kolom dan penggabungan: lihat `tokos_from_rows`."""
    if not path.is_file():
        raise StoreListMissing(
            f"Daftar toko belum tersedia: letakkan file '{path.name}' di folder '{path.parent.name}/' "
            f"dengan kolom {list(STORE_HEADERS_HELP)}.")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ReferenceDataError(f"File '{path.name}' tidak dapat dibaca: {exc}") from None
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        return tokos_from_rows(ws.iter_rows(values_only=True), path.name, row_filter=row_filter, max_listed=max_listed)
    finally:
        wb.close()


def load_stores_from_values(values: list[list], source: str, *, row_filter: RowFilter | None = None) -> list[Toko]:
    """Daftar toko dari isi Google Sheet (`worksheet.get_all_values()`)."""
    return tokos_from_rows(values, source, row_filter=row_filter)


def _find_store_columns(header: list[str]):
    """(indeks region, indeks grup | None, indeks toko) bila header lengkap, selain itu None."""
    def find(aliases):
        return next((header.index(a) for a in aliases if a in header), None)
    region, toko = find(REGION_COLS), find(TOKO_COLS)
    if region is None or toko is None:
        return None
    return region, find(GROUP_COLS), toko


def _header_matches(cells) -> bool:
    cells = tuple(_clean(c) for c in cells[: len(EXPECTED_HEADERS)])
    return len(cells) == len(EXPECTED_HEADERS) and cells[1:] == EXPECTED_HEADERS[1:] and \
        cells[0] in (EXPECTED_HEADERS[0], LEGACY_FIRST_HEADER)


def load_catalogue(path: Path) -> list[Sku]:
    """Daftar SKU + nama produk + isi + nomor BPOM resmi, dibaca dari Sheet1 template resmi dari BD.

    Template adalah sumber tunggal: yang di-download pengguna dan yang dipakai validasi selalu sama.
    Header kolom A boleh "Nama Toko (By Group)" atau "Distributor Branch" (template lama).
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ReferenceDataError(f"Template '{path.name}' tidak dapat dibaca: {exc}") from None
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        rows = enumerate(ws.iter_rows(values_only=True), start=1)
        header_no = None
        for no, cells in rows:
            if no > HEADER_SEARCH_ROWS:
                break
            if _header_matches(cells):
                header_no = no
                break
        if header_no is None:
            raise ReferenceDataError(f"Header template '{path.name}' harus {list(EXPECTED_HEADERS)}.")
        result, seen = [], set()
        for no, cells in rows:
            cells = tuple(cells) + (None,) * (len(EXPECTED_HEADERS) - len(cells))
            _, product_id, product_name, isi, bpom, _ = (_clean(c) for c in cells[: len(EXPECTED_HEADERS)])
            if not (product_id or product_name or isi or bpom):
                continue
            if not (product_id and product_name and isi and bpom):
                raise ReferenceDataError(
                    f"'{path.name}' baris {no}: PRODUCT ID, PRODUCT NAME, ISI, dan BPOM NUMBER wajib diisi.")
            key = row_key(product_id, bpom)
            if key in seen:
                raise ReferenceDataError(f"'{path.name}' baris {no}: kombinasi {product_id} + {bpom} ganda.")
            seen.add(key)
            result.append(Sku(product_id, product_name, bpom, isi))
    finally:
        wb.close()
    if not result:
        raise ReferenceDataError(f"'{path.name}' tidak berisi SKU.")
    return result
