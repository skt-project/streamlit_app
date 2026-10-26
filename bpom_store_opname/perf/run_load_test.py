"""Uji beban end-to-end: N pengguna (browser sungguhan) menjalani alur lengkap ke server Streamlit ASLI.

Alur tiap pengguna: buka halaman -> pilih Region -> pilih Toko -> tanggal + PIC -> (sebagian) tambah SKU lewat form -> download
template -> isi QTY -> upload -> periksa rincian final -> centang konfirmasi -> Submit -> "Berhasil menyimpan".

Server memakai aplikasi asli (app.main) dengan Google Sheets TIRUAN berkuota asli (60 baca + 60 tulis per menit) dan latensi
(perf/fakegoogle.py): tidak ada data uji yang menyentuh spreadsheet asli.

Jalankan dengan Python yang punya Playwright (+ chromium) dan openpyxl; server memakai .venv proyek:
    python perf/run_load_test.py --users 100 --ramp 5

Laporan memisahkan BEBAN SERVER (CPU/memori/Sheets/integritas data) dari LATENSI KLIEN, karena browser penguji berbagi mesin
dengan server: bila CPU mesin jenuh, latensi klien dibatasi mesin penguji, bukan aplikasi.
"""
from __future__ import annotations

import argparse
import asyncio
import io
import json
import math
import os
import statistics
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
REGIONS = [f"Region {n}" for n in range(1, 8)]
GROUPS_PER_REGION = 2
CATALOGUE_ROWS = 66


# ── data uji ────────────────────────────────────────────────────────────────

def build_stores(path: Path, users: int) -> int:
    """Buat daftar toko buatan; mengembalikan jumlah pilihan per region (>= pengguna per region, supaya tiap pengguna unik)."""
    import openpyxl
    per_region = math.ceil(users / len(REGIONS)) + 2
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Region", "Nama Group Toko", "Nama Toko"])
    for region in REGIONS:
        n = region.split()[-1]
        for g in range(1, GROUPS_PER_REGION + 1):                       # satu grup (3 cabang) = SATU pilihan
            for c in range(3):
                ws.append([region, f"GRUP {n}{g}", f"Cabang {g}-{c}"])
        for k in range(per_region - GROUPS_PER_REGION):
            ws.append([region, None, f"TOKO {n}-{k:03d}"])
    wb.save(path)
    return per_region


def stats(values: list[float]) -> dict | None:
    if not values:
        return None
    values = sorted(values)
    p95 = values[min(len(values) - 1, math.ceil(0.95 * len(values)) - 1)]
    return {"n": len(values), "median": round(statistics.median(values), 2), "p95": round(p95, 2), "max": round(values[-1], 2)}


# ── server ──────────────────────────────────────────────────────────────────

def _health(port: int) -> bool:
    try:
        with urllib.request.urlopen(f"http://localhost:{port}/_stcore/health", timeout=2) as r:
            return r.status == 200
    except Exception:
        return False


def pin_cpus(pid: int, mask: int) -> None:
    """Batasi proses `pid` ke inti CPU tertentu (bitmask). Anak proses yang dibuat sesudahnya mewarisi pembatasan ini."""
    subprocess.run(["powershell", "-NoProfile", "-Command", f"(Get-Process -Id {pid}).ProcessorAffinity = {mask}"],
                   capture_output=True, timeout=30)


def affinity_masks(server_cores: int) -> tuple[int, int]:
    """(mask server, mask klien): server di `server_cores` inti pertama, browser/driver di inti sisanya."""
    total = os.cpu_count() or 1
    if not 0 < server_cores < total:
        raise SystemExit(f"--server-cores harus 1..{total - 1} (mesin ini punya {total} inti logis).")
    server = (1 << server_cores) - 1
    return server, ((1 << total) - 1) & ~server


def start_server(port: int, out: Path, stores: Path, venv_python: Path) -> subprocess.Popen:
    if _health(port):
        raise SystemExit(f"Port {port} sudah dipakai server lain; hentikan dulu atau pilih --port lain.")
    import os
    env = {**os.environ, "PERF_OUT": str(out), "PERF_STORES": str(stores), "PYTHONUNBUFFERED": "1"}
    proc = subprocess.Popen(
        [str(venv_python), "-m", "streamlit", "run", str(HERE / "harness.py"), "--server.port", str(port),
         "--server.headless", "true", "--browser.gatherUsageStats", "false"],
        cwd=ROOT, env=env, stdout=open(out / "server.log", "w", encoding="utf-8"), stderr=subprocess.STDOUT)
    for _ in range(60):
        if _health(port):
            return proc
        if proc.poll() is not None:
            raise SystemExit(f"Server berhenti saat start; lihat {out / 'server.log'}")
        time.sleep(1)
    proc.terminate()
    raise SystemExit("Server tidak siap dalam 60 detik.")


def listener_pid(port: int) -> int | None:
    """PID proses yang BENAR-BENAR melayani port (python.exe di venv Windows hanyalah peluncur yang memunculkan proses anak)."""
    script = f"(Get-NetTCPConnection -LocalPort {port} -State Listen -ErrorAction SilentlyContinue | Select-Object -First 1).OwningProcess"
    try:
        return int(subprocess.run(["powershell", "-NoProfile", "-Command", script], capture_output=True, text=True, timeout=30).stdout.strip())
    except Exception:
        return None


def _sample(pid: int) -> dict | None:
    script = (f"$p = Get-Process -Id {pid} -ErrorAction SilentlyContinue; if ($p) {{ '{{0}} {{1}} {{2}}' -f $p.WorkingSet64, $p.CPU, "
              "((Get-CimInstance Win32_Processor | Measure-Object LoadPercentage -Average).Average) }")
    try:
        text = subprocess.run(["powershell", "-NoProfile", "-Command", script], capture_output=True, text=True, timeout=30).stdout.split()
        rss, cpu, load = float(text[0]), float(text[1]), float(text[2])
        return {"t": time.time(), "rss_mb": rss / 1048576, "cpu_s": cpu, "machine_cpu_pct": load}
    except Exception:
        return None


class Sampler(threading.Thread):
    def __init__(self, pid: int, every: float = 5.0):
        super().__init__(daemon=True)
        self.pid, self.every, self.samples, self._halt = pid, every, [], threading.Event()

    def run(self):
        while not self._halt.is_set():
            s = _sample(self.pid)
            if s:
                self.samples.append(s)
            self._halt.wait(self.every)

    def stop(self):
        self._halt.set()


# ── pengguna ────────────────────────────────────────────────────────────────

async def user_journey(i: int, browser, url: str, workdir: Path, extra_share: float, timings: dict, failures: list, outcomes: list,
                       load_only: bool = False):
    import openpyxl
    region_idx, toko_idx = i % len(REGIONS), i // len(REGIONS)
    with_extra = (i % 100) < int(extra_share * 100)
    step = "start"
    ctx = await browser.new_context(viewport={"width": 375, "height": 667}, accept_downloads=True)
    pg = await ctx.new_page()
    pg.set_default_timeout(240_000)
    t_all = time.time()

    async def pick(label: str, nth: int):
        """Buka dropdown lalu pilih opsi ke-nth. Seperti pengguna sungguhan: bila dropdown menutup sendiri karena halaman
        me-render ulang saat server sibuk, klik lagi (maks. 4x). Setiap pembukaan ulang dicatat di laporan."""
        from playwright.async_api import TimeoutError as PWTimeout
        for attempt in range(4):
            await pg.get_by_role("combobox", name=label, exact=True).click(timeout=60_000)
            await pg.keyboard.press("ArrowDown")
            try:
                await pg.wait_for_selector('[role="option"]', timeout=20_000)
                break
            except PWTimeout:
                timings.setdefault("dropdown_dibuka_ulang", []).append(float(attempt + 1))
                await pg.keyboard.press("Escape")
        else:
            raise AssertionError(f"dropdown '{label}' tidak terbuka setelah 4 percobaan")
        await pg.locator('[role="option"]').nth(nth).click(timeout=60_000)

    def mark(name: str, t0: float):
        timings.setdefault(name, []).append(time.time() - t0)

    try:
        step, t = "load", time.time()
        await pg.goto(url, timeout=240_000)
        await pg.wait_for_selector("text=Langkah 1")
        mark("load", t)
        if load_only:                                                          # hanya mengukur biaya membuka halaman
            outcomes.append({"user": i, "expected_rows": 0})
            return

        step, t = "pilih_region_toko", time.time()
        await pick("Region", region_idx)
        await pg.wait_for_function("(() => {const e = document.querySelector('input[aria-label=Toko]'); return e && !e.disabled})()")
        await pick("Toko", toko_idx)
        mark("pilih_region_toko", t)

        step, t = "tanggal_pic", time.time()
        await pg.get_by_test_id("stDateInput").get_by_text("dd").first.click()
        await pg.keyboard.type("26102026")
        await pg.keyboard.press("Enter")
        await pg.get_by_role("textbox", name="Nama PIC").fill(f"PIC Uji {i}")
        await pg.keyboard.press("Tab")
        await pg.wait_for_timeout(800)
        mark("tanggal_pic", t)

        if with_extra:
            step, t = "tambah_sku", time.time()
            await pick("SKU", 0)
            await pg.get_by_role("textbox", name="Nomor BPOM").fill(f"NA{i:011d}")
            await pg.get_by_role("button", name="Tambah ke daftar").click()
            await pg.wait_for_selector("text=Daftar tambahan (1 baris)")
            mark("tambah_sku", t)

        step, t = "download_template", time.time()
        async with pg.expect_download(timeout=240_000) as dl:
            await pg.get_by_role("button", name="DOWNLOAD TEMPLATE").click()
        path = workdir / f"dl_{i}.xlsx"
        await (await dl.value).save_as(str(path))
        mark("download_template", t)
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        for r in range(3, ws.max_row + 1):
            ws.cell(r, 6).value = (r + i) % 9                                   # QTY 0..8 (termasuk 0)
        buf = io.BytesIO()
        wb.save(buf)

        step, t = "upload_valid", time.time()
        await pg.locator('input[type="file"]').set_input_files(
            {"name": "isi.xlsx", "mimeType": "application/octet-stream", "buffer": buf.getvalue()})
        await pg.wait_for_selector("text=baris siap disimpan")
        mark("upload_valid", t)

        step, t = "rincian_final", time.time()
        expected = CATALOGUE_ROWS + (1 if with_extra else 0)
        summary = pg.locator("li", has_text="Total baris yang akan disimpan").first      # seluruh butir, bukan hanya teks tebalnya
        await summary.wait_for()
        text = await summary.inner_text()
        if f": {expected} (" not in text:
            raise AssertionError(f"rincian final salah: harapan {expected} baris, tampil '{text}'")
        await pg.locator("table").filter(has_text="SKU - Product Name").first.wait_for()
        mark("rincian_final", t)

        step, t = "submit", time.time()
        await pg.get_by_text("Saya sudah memeriksa rincian").click()
        await pg.get_by_role("button", name="Submit").click()
        await pg.wait_for_selector("text=Berhasil menyimpan")
        mark("submit", t)
        timings.setdefault("total_journey", []).append(time.time() - t_all)
        outcomes.append({"user": i, "expected_rows": expected})
    except Exception as exc:                                                   # noqa: BLE001
        lines = str(exc).splitlines()
        failures.append({"user": i, "step": step, "error": f"{type(exc).__name__}: {lines[0][:140]}",
                         "detail": " | ".join(x.strip() for x in lines[1:8])[:500]})        # log panggilan Playwright (mis. 'intercepts pointer events')
        try:                                                                   # bukti visual: bagaimana layar saat gagal
            await pg.screenshot(path=str(workdir / f"fail_{i}_{step}.png"), timeout=20_000)
            (workdir / f"fail_{i}_{step}.txt").write_text(await pg.evaluate("document.body.innerText"), encoding="utf-8")
        except Exception:                                                      # noqa: BLE001
            pass
    finally:
        await ctx.close()


async def run_users(args, url: str, workdir: Path):
    from playwright.async_api import async_playwright
    timings: dict = {}
    failures: list = []
    outcomes: list = []
    async with async_playwright() as p:
        browsers = [await p.chromium.launch(args=["--disable-gpu", "--disable-dev-shm-usage"])
                    for _ in range(math.ceil(args.users / args.per_browser))]
        tasks = []
        t0 = time.time()
        for i in range(args.users):
            await asyncio.sleep(1.0 / args.ramp if i else 0)
            tasks.append(asyncio.create_task(user_journey(
                i, browsers[i // args.per_browser], url, workdir, args.extra_sku_share, timings, failures, outcomes, args.load_only)))
        await asyncio.gather(*tasks)
        wall = time.time() - t0
        for b in browsers:
            await b.close()
    return timings, failures, outcomes, wall


# ── laporan ─────────────────────────────────────────────────────────────────

def build_report(args, per_region, timings, failures, outcomes, wall, samples, out: Path, t_start: float) -> dict:
    calls = json.loads((out / "calls.json").read_text(encoding="utf-8")) if (out / "calls.json").exists() else {}
    runs = []
    if (out / "runs.log").exists():
        runs = [tuple(map(float, line.split())) for line in (out / "runs.log").read_text().splitlines() if line.strip()]
    walls, cpus = [r[0] for r in runs], [r[1] for r in runs]
    rows = [tuple(r) for r in calls.get("rows", [])]
    expected_rows = sum(o["expected_rows"] for o in outcomes)
    per_toko: dict = {}
    for r in rows:
        per_toko.setdefault((r[0], r[1]), 0)
        per_toko[(r[0], r[1])] += 1
    # Biaya CPU server SELAMA beban saja: dikurangi sampel terakhir sebelum pengguna pertama (tanpa biaya start/impor awal).
    before = [x for x in samples if x["t"] <= t_start]
    baseline = before[-1] if before else (samples[0] if samples else None)
    cpu_delta = (samples[-1]["cpu_s"] - baseline["cpu_s"]) if baseline and len(samples) > 1 else None
    idle_rate = ((before[-1]["cpu_s"] - before[0]["cpu_s"]) / (before[-1]["t"] - before[0]["t"])) if len(before) > 1 else 0.0
    done = len(outcomes)
    report = {
        "config": {"users": args.users, "ramp_users_per_s": args.ramp, "extra_sku_share": args.extra_sku_share,
                   "viewport": "375x667", "options_per_region": per_region,
                   "server_cores": args.server_cores or f"tanpa batas ({os.cpu_count()} inti logis dibagi dengan browser)"},
        "result": {"completed": done, "failed": len(failures), "wall_s": round(wall, 1),
                   "failures_by_step": {s: sum(1 for f in failures if f["step"] == s) for s in sorted({f["step"] for f in failures})},
                   "sample_failures": failures[:6]},
        "client_latency_s": {k: stats(v) for k, v in timings.items()},
        "server": {
            "rss_mb_before_load": round(baseline["rss_mb"]) if baseline else None,
            "rss_mb_peak": round(max(s["rss_mb"] for s in samples)) if samples else None,
            "cpu_seconds_during_load": round(cpu_delta, 1) if cpu_delta is not None else None,
            "cpu_idle_rate_pct_of_core": round(idle_rate * 100, 1),
            "cpu_seconds_per_completed_user": round(cpu_delta / done, 3) if cpu_delta is not None and done else None,
            "script_runs": len(runs),
            "script_wall_s": stats(walls), "script_thread_cpu_ms": {k: round(v * 1000, 1) for k, v in (stats(cpus) or {}).items() if k != "n"},
            "machine_cpu_pct_avg": round(statistics.mean(s["machine_cpu_pct"] for s in samples)) if samples else None,
            "machine_cpu_pct_max": round(max(s["machine_cpu_pct"] for s in samples)) if samples else None,
        },
        "sheets": {
            "reads": calls.get("reads"), "batch_updates": calls.get("writes"), "http_429_read": calls.get("429_read"),
            "http_429_write": calls.get("429_write"), "spreadsheet_opens": calls.get("opened"), "worksheet_ensures": calls.get("ensure"),
            "submits_per_batch": calls.get("requests_per_call"), "max_concurrent_writes": calls.get("max_inflight_writes"),
            "mutation_kinds": sorted(set(calls.get("mutations", []))),
        },
        "integrity": {
            "rows_written": len(rows), "rows_expected_from_completed_users": expected_rows,
            "duplicate_rows": len(rows) - len(set(rows)), "tokos_written": len(per_toko),
            "rows_per_toko_values": sorted(set(per_toko.values())),
        },
    }
    srv, sh, it = report["server"], report["sheets"], report["integrity"]
    window_model = {}
    if srv["cpu_seconds_per_completed_user"]:
        for w in (30, 60, 120):
            window_model[f"100 pengguna dalam {w}s"] = f"{100 * srv['cpu_seconds_per_completed_user'] / w * 100:.0f}% dari satu inti"
        window_model["catatan"] = ("CPU proses server selama beban dibagi pengguna selesai; dengan <30 pengguna biaya tetap "
                                   "(koneksi, file statis) membuatnya terlalu tinggi")
    report["capacity_model_server_cpu"] = window_model
    checks = {
        "tanpa HTTP 429 dari Sheets": (sh["http_429_read"] or 0) + (sh["http_429_write"] or 0) == 0,
        "hanya appendCells yang menulis": sh["mutation_kinds"] in ([], ["appendCells"]),
        "tanpa baris ganda": it["duplicate_rows"] == 0,
        "baris tersimpan == harapan": it["rows_written"] == expected_rows,
        "memori server puncak < 700 MB": (srv["rss_mb_peak"] or 0) < 700,
        "CPU eksekusi script p95 < 300 ms": ((srv["script_thread_cpu_ms"] or {}).get("p95") or 0) < 300,
        "spreadsheet dibuka hanya sekali": sh["spreadsheet_opens"] in (1, None, 0),
    }
    report["checks"] = checks
    report["client_note"] = ("CPU mesin penguji jenuh (max {}%): latensi klien dibatasi laptop penguji, bukan aplikasi.".format(srv["machine_cpu_pct_max"])
                             if (srv["machine_cpu_pct_max"] or 0) >= 85 else "CPU mesin penguji tidak jenuh: latensi klien mendekati kondisi nyata.")
    return report


def print_report(rep: dict) -> None:
    c, r = rep["config"], rep["result"]
    print(f"\n=== UJI BEBAN {c['users']} pengguna (ramp {c['ramp_users_per_s']}/dtk, layar {c['viewport']}, server: {c['server_cores']} inti) ===")
    print(f"Selesai: {r['completed']}/{c['users']}  gagal: {r['failed']}  total waktu: {r['wall_s']} dtk")
    if r["failures_by_step"]:
        print("Gagal per langkah:", r["failures_by_step"])
        for f in r["sample_failures"]:
            print("   -", f)
    print("\nLatensi klien (dtk):")
    for k, v in rep["client_latency_s"].items():
        print(f"  {k:20s} {v}")
    s = rep["server"]
    print(f"\nServer: memori {s['rss_mb_before_load']} -> puncak {s['rss_mb_peak']} MB | CPU selama beban {s['cpu_seconds_during_load']} dtk "
          f"({s['cpu_seconds_per_completed_user']} dtk/pengguna) | CPU saat idle {s['cpu_idle_rate_pct_of_core']}% inti")
    print(f"  eksekusi script: {s['script_runs']}x, durasi dinding {s['script_wall_s']}, CPU thread (ms) {s['script_thread_cpu_ms']}")
    print(f"  CPU mesin: rata-rata {s['machine_cpu_pct_avg']}%, maks {s['machine_cpu_pct_max']}%")
    print(f"Sheets: {rep['sheets']}")
    print(f"Integritas: {rep['integrity']}")
    print("Model kapasitas (CPU server, 1 inti efektif karena GIL):", rep["capacity_model_server_cpu"])
    print("\nPemeriksaan:")
    for name, ok in rep["checks"].items():
        print(f"  [{'LULUS' if ok else 'GAGAL'}] {name}")
    print("\nCatatan:", rep["client_note"])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--users", type=int, default=100)
    ap.add_argument("--ramp", type=float, default=5.0, help="pengguna baru per detik (default 5; 100 pengguna masuk dalam 20 dtk)")
    ap.add_argument("--port", type=int, default=8599)
    ap.add_argument("--per-browser", type=int, default=10, help="konteks per proses Chromium")
    ap.add_argument("--extra-sku-share", type=float, default=0.3, help="porsi pengguna yang menambah SKU lewat form")
    ap.add_argument("--server-cores", type=int, default=0,
                    help="kunci server ke N inti CPU pertama dan browser ke inti sisanya (meniru wadah cloud kecil, dan mencegah "
                         "browser penguji merebut CPU server). 0 = tanpa pembatasan")
    ap.add_argument("--pyspy", type=int, default=0, metavar="DETIK",
                    help="rekam profil server dengan py-spy selama DETIK dtk (pip install py-spy di venv proyek); ringkas dengan summarize_profile.py")
    ap.add_argument("--load-only", action="store_true", help="sesi hanya membuka halaman (mengukur biaya memuat halaman saja)")
    ap.add_argument("--venv-python", type=Path, default=ROOT / ".venv" / "Scripts" / "python.exe")
    ap.add_argument("--out", type=Path, default=None, help="folder keluaran (default: folder sementara)")
    args = ap.parse_args()

    out = args.out or Path(tempfile.mkdtemp(prefix="bpom_perf_"))
    out.mkdir(parents=True, exist_ok=True)
    stores = out / "stores.xlsx"
    per_region = build_stores(stores, args.users)
    print(f"Keluaran: {out}\nDaftar toko buatan: {len(REGIONS)} region x {per_region} pilihan")
    proc = start_server(args.port, out, stores, args.venv_python)
    server_pid = listener_pid(args.port) or proc.pid
    print(f"Server: peluncur pid {proc.pid}, proses server pid {server_pid}")
    if args.server_cores:
        server_mask, client_mask = affinity_masks(args.server_cores)
        pin_cpus(server_pid, server_mask)
        pin_cpus(os.getpid(), client_mask)                                     # browser yang dibuka driver mewarisi ini
        print(f"CPU dikunci: server {args.server_cores} inti (mask {server_mask:#x}), driver+browser mask {client_mask:#x}")
    sampler = Sampler(server_pid)
    sampler.start()
    time.sleep(6)                                                              # memori/CPU dasar sebelum beban
    profiler = None
    if args.pyspy:                                                             # profil: siapa yang memegang GIL server selama beban
        spy = args.venv_python.parent / "py-spy.exe"
        profiler = subprocess.Popen(
            [str(spy), "record", "-p", str(server_pid), "-d", str(args.pyspy), "-r", "50", "-f", "raw", "--threads", "--gil",
             "--nonblocking", "-o", str(out / "profile.raw")], stdout=subprocess.DEVNULL, stderr=open(out / "pyspy.log", "w"))
        print(f"py-spy merekam {args.pyspy} dtk ke {out / 'profile.raw'}")
    t_start = time.time()
    try:
        timings, failures, outcomes, wall = asyncio.run(run_users(args, f"http://localhost:{args.port}", out))
        time.sleep(8)                                                          # biarkan penulisan terakhir selesai
    finally:
        sampler.stop()
        sampler.join(timeout=35)
        subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True)    # peluncur + proses anaknya
        if profiler:
            try:
                profiler.wait(timeout=90)                                      # py-spy menulis berkasnya saat target berhenti/durasi habis
            except subprocess.TimeoutExpired:
                profiler.kill()
    report = build_report(args, per_region, timings, failures, outcomes, wall, sampler.samples, out, t_start)
    (out / "report.json").write_text(json.dumps(report, indent=1, ensure_ascii=False), encoding="utf-8")
    print_report(report)
    print(f"\nLaporan lengkap: {out / 'report.json'}")
    return 0 if all(report["checks"].values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
