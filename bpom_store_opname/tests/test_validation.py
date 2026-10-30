"""Tes validasi, template, dan penyimpanan: python -m pytest tests"""
import sys
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import reference_data as r  # noqa: E402
import sheets_store as store  # noqa: E402
import validation as v  # noqa: E402
from template_builder import build_template  # noqa: E402

NAME = "CV ANDIRA DAHAYU - BENGKULU"
NAME2 = "CV ANDIRA DAHAYU - CURUP"
TEMPLATE = ROOT / "template" / "Contoh Excel Response.xlsx"


def make(rows, headers=v.EXPECTED_HEADERS, title=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    if title:
        ws.append([None] * (len(v.EXPECTED_HEADERS) - 1) + ["Filled by Distributor"])
    ws.append(list(headers))
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def good(qty=5, pid="P-1", bpom="NA1", name=NAME, isi="10g"):
    return [name, pid, "Produk", isi, bpom, qty]


def messages(res):
    return " | ".join(i.message for i in res.issues)


def catalogue():
    return r.load_catalogue(TEMPLATE)


# ── parsing & aturan kolom ──────────────────────────────────────────────────

def test_valid_file_with_title_row():
    res = v.parse_excel(make([good(), good(qty=0, bpom="NA2"), good(qty=3, bpom="NA3")]), [NAME])
    assert res.ok and [x.qty for x in res.rows] == [5, 0, 3]
    assert res.rows[0].excel_row == 3 and res.rows[0].isi == "10g"


def test_name_match_ignores_case_and_spacing():
    assert v.parse_excel(make([good(name="cv andira dahayu-bengkulu")]), [NAME]).ok


def test_header_exact_case_sensitive_and_new_name():
    h = list(v.EXPECTED_HEADERS); h[5] = "Qty"
    res = v.parse_excel(make([good()], headers=h), [NAME])
    assert not res.ok and "persis" in messages(res)
    old = ["Distributor Branch", *v.EXPECTED_HEADERS[1:]]                       # header lama ("Distributor Branch") ditolak saat upload
    assert not v.parse_excel(make([good()], headers=old), [NAME]).ok
    assert v.EXPECTED_HEADERS[0] == "Nama Toko (By Group)"


def test_extra_missing_and_reordered_columns():
    assert not v.parse_excel(make([good() + ["x"]], headers=[*v.EXPECTED_HEADERS, "EXTRA"]), [NAME]).ok
    assert not v.parse_excel(make([good()[:5]], headers=v.EXPECTED_HEADERS[:5]), [NAME]).ok
    swapped = [v.EXPECTED_HEADERS[1], v.EXPECTED_HEADERS[0], *v.EXPECTED_HEADERS[2:]]
    assert "Urutan" in messages(v.parse_excel(make([good()], headers=swapped), [NAME]))


def test_qty_rules_whole_non_negative():
    for bad in (None, "", 1.5, "abc", "1,5", True, -1, -12.0, "-12", -2**40):
        assert v.parse_qty(bad)[0] is None, bad
    assert "negatif" in v.parse_qty(-3)[1]
    assert v.parse_qty(7.0) == (7, None)
    assert v.parse_qty("12") == (12, None)
    assert v.parse_qty(0) == (0, None)
    assert v.parse_qty("-0") == (0, None)


def test_negative_qty_row_is_rejected_with_row_number():
    res = v.parse_excel(make([good(), good(qty=-3, bpom="NA2")]), [NAME])
    assert not res.ok and [(i.row, i.column) for i in res.issues] == [(4, "QTY")] and "negatif" in messages(res)


def test_all_or_nothing_and_row_numbers():
    res = v.parse_excel(make([good(), good(qty=None, bpom="NA2"), good(bpom="NA3")]), [NAME])
    assert not res.ok and [i.row for i in res.issues] == [4]


def test_duplicate_key_within_file():
    res = v.parse_excel(make([good(), good(pid="p-1", bpom="na1")]), [NAME])
    assert not res.ok and "Duplikat dengan baris 3" in messages(res)


def test_required_text_fields_including_isi():
    res = v.parse_excel(make([["CV LAIN - X", "", "N", "", "B", 1]]), [NAME])
    cols = {i.column for i in res.issues}
    assert {v.EXPECTED_HEADERS[0], "PRODUCT ID", v.COL_ISI} <= cols


def test_blank_rows_skipped_empty_and_non_xlsx_rejected():
    assert v.parse_excel(make([good(), [None] * 6, [None] * 6]), [NAME]).ok
    assert not v.parse_excel(make([]), [NAME]).ok
    assert not v.parse_excel(b"bukan excel", [NAME]).ok


def test_form_validation():
    assert v.validate_form(region="A", toko="S", tanggal="d", pic="x", has_file=True) == []
    assert len(v.validate_form(region=None, toko=None, tanggal=None, pic="", has_file=False)) == 5


# ── aturan SKU / BPOM tambahan / multi distributor ──────────────────────────

SKUS = [v.Sku("P-1", "Produk", "NA1", "10g"), v.Sku("P-1", "Produk", "NA2", "10g"), v.Sku("P-2", "Lain", "NB1", "5ml")]


def test_sku_rules_known_additional_and_missing():
    rows = [good(), good(bpom="NA2"), good(bpom="NA-BARU", qty=3)]
    res = v.parse_excel(make(rows), [NAME], SKUS)
    assert res.ok and [x.additional for x in res.rows] == [False, False, True]
    assert [k.product_id for _, k in res.missing] == ["P-2"]


def test_unknown_sku_name_and_isi_mismatch_rejected():
    res = v.parse_excel(make([good(pid="P-9")]), [NAME], SKUS)
    assert not res.ok and any(i.column == "PRODUCT ID" for i in res.issues)
    bad_name = good(); bad_name[2] = "Nama Lain"
    assert any(i.column == "PRODUCT NAME" for i in v.parse_excel(make([bad_name]), [NAME], SKUS).issues)
    res = v.parse_excel(make([good(isi="99ml")]), [NAME], SKUS)
    assert not res.ok and any(i.column == v.COL_ISI for i in res.issues)
    assert v.parse_excel(make([good(isi="10G")]), [NAME], SKUS).ok      # beda huruf besar/kecil tidak masalah


def test_distributor_must_belong_to_company_and_same_sku_allowed_per_distributor():
    res = v.parse_excel(make([good(name="CV LAIN - X")]), [NAME, NAME2])
    assert not res.ok and "bukan toko yang dipilih" in messages(res)
    assert v.parse_excel(make([good(), good(name=NAME2, qty=2)]), [NAME, NAME2]).ok
    rows = [good(), good(name=NAME2, qty=2), good(name=NAME2.lower(), pid="p-1", bpom="na1", qty=3)]
    assert "Duplikat" in messages(v.parse_excel(make(rows), [NAME, NAME2]))


def test_missing_is_reported_per_distributor():
    res = v.parse_excel(make([good()]), [NAME, NAME2], [v.Sku("P-1", "Produk", "NA1", "10g")])
    assert res.ok and [(n, k.bpom_number) for n, k in res.missing] == [(NAME2, "NA1")]


# ── data referensi & template ───────────────────────────────────────────────

TOKO_LABEL = "TOKO MAJU JAYA"
REGION = "SUMATERA"


def _xlsx(tmp_path, name, rows):
    wb = openpyxl.Workbook(); ws = wb.active
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


def test_catalogue_loads_from_current_template():
    skus = catalogue()
    assert len(skus) == 66 and all(s.isi and s.product_name for s in skus)


def test_load_stores_region_group_and_toko_columns_any_order_with_title_rows(tmp_path):
    path = _xlsx(tmp_path, "s.xlsx", [
        ["DAFTAR TOKO"], [None],
        ["Nama Toko", "Catatan", "NAMA GROUP TOKO", "Region"],                 # urutan bebas, huruf bebas, kolom lain diabaikan
        ["Alfa Cabang 1", "-", "ALFAMART", "Sumatera"],
        ["Alfa Cabang 2", "-", "ALFAMART", "Sumatera"],                         # satu grup = SATU pilihan
        ["Alfa Cabang 3", "-", "ALFAMART", "Jawa"],                             # grup yang sama di region lain = pilihan lain
        ["Toko Budi", None, None, "Sumatera"],                                  # tanpa grup -> nama toko
        [None, None, None, None],
    ])
    tokos = r.load_stores(path)
    assert [(t.region, t.name, t.is_group, t.outlets) for t in tokos] == [
        ("Sumatera", "ALFAMART", True, 2), ("Jawa", "ALFAMART", True, 1), ("Sumatera", "Toko Budi", False, 1)]
    assert tokos[0].label == "ALFAMART  ·  grup (2 toko)" and tokos[2].label == "Toko Budi"


def test_load_stores_works_without_group_column(tmp_path):
    path = _xlsx(tmp_path, "s.xlsx", [["Region", "Nama Toko"], ["Jawa", "Toko A"], ["jawa", "TOKO A"]])
    (t,) = r.load_stores(path)
    assert (t.name, t.outlets) == ("Toko A", 2)                                 # nama sama di region sama = satu pilihan


def test_load_stores_rejects_bad_files(tmp_path):
    import pytest
    with pytest.raises(r.StoreListMissing, match="belum tersedia"):
        r.load_stores(tmp_path / "tidak-ada.xlsx")
    with pytest.raises(r.ReferenceDataError, match="Header"):
        r.load_stores(_xlsx(tmp_path, "a.xlsx", [["Nama Toko"], ["x"]]))                     # tanpa Region
    head = ["Region", "Nama Group Toko", "Nama Toko"]
    with pytest.raises(r.ReferenceDataError, match="wajib diisi") as exc:
        r.load_stores(_xlsx(tmp_path, "b.xlsx", [head, ["", None, "A"], ["Jawa", None, ""], ["Jawa", "G", "B"]]))
    assert "baris 2" in str(exc.value) and "baris 3" in str(exc.value)                      # semua masalah, bukan satu
    with pytest.raises(r.ReferenceDataError, match="grup dan sebagai toko tunggal"):
        r.load_stores(_xlsx(tmp_path, "c.xlsx", [head, ["Jawa", "ALFA", "x"], ["Jawa", None, "ALFA"]]))
    with pytest.raises(r.ReferenceDataError, match="tidak berisi toko"):
        r.load_stores(_xlsx(tmp_path, "d.xlsx", [head]))


def test_template_is_for_one_toko_and_original_untouched():
    skus = catalogue()
    raw = TEMPLATE.read_bytes()
    out = build_template(raw, REGION, TOKO_LABEL, skus)
    assert raw == TEMPLATE.read_bytes()                         # template asli tidak berubah
    wb = openpyxl.load_workbook(BytesIO(out))
    assert wb.sheetnames == ["Sheet1", v.INSTRUCTION_SHEET]      # sheet kerja BD tidak ikut
    ws = wb["Sheet1"]
    assert [c.value for c in ws[2]] == list(v.EXPECTED_HEADERS) and ws["F1"].value == "Filled by Distributor"
    assert ws["A2"].value == "Nama Toko (By Group)"             # template lama "Distributor Branch" -> header toko
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=3)]
    assert rows == [[TOKO_LABEL, k.product_id, k.product_name, k.isi, k.bpom_number, None] for k in skus]
    (dv,) = ws.data_validations.dataValidation                  # QTY: bilangan bulat >= 0 di Excel
    assert (dv.type, dv.operator, dv.formula1) == ("whole", "greaterThanOrEqual", "0")
    assert str(dv.sqref) == f"F3:F{ws.max_row}"
    guide = " ".join(str(c.value) for c in wb[v.INSTRUCTION_SHEET]["A"] if c.value)
    assert TOKO_LABEL in guide and REGION in guide and "Satu file = satu toko" in guide and "26 Oktober" in guide
    assert "Distributor Branch" not in guide and "23 Oktober" not in guide


def test_filled_template_passes_and_ignores_guide_sheet():
    skus = catalogue()
    wb = openpyxl.load_workbook(BytesIO(build_template(TEMPLATE.read_bytes(), REGION, TOKO_LABEL, skus)))
    ws = wb["Sheet1"]
    for i in range(3, ws.max_row + 1):
        ws.cell(i, 6).value = i
    buf = BytesIO(); wb.save(buf)
    res = v.parse_excel(buf.getvalue(), [TOKO_LABEL], skus)
    assert res.ok and len(res.rows) == len(skus) and not res.missing and res.warnings == []
    other = v.parse_excel(buf.getvalue(), ["TOKO LAIN"], skus)                  # file milik toko lain ditolak
    assert not other.ok and "bukan toko yang dipilih" in messages(other)


def test_catalogue_still_loads_when_template_has_new_header(tmp_path):
    out = build_template(TEMPLATE.read_bytes(), REGION, TOKO_LABEL, catalogue())      # header baru "Nama Toko (By Group)"
    path = tmp_path / "t.xlsx"; path.write_bytes(out)
    assert len(r.load_catalogue(path)) == 66


# ── Google Sheets (mock) ────────────────────────────────────────────────────

def _ws(lines=()):
    """Worksheet palsu; `lines` = baris yang sudah ada di sheet ({header: nilai})."""
    ws = MagicMock(); ws.id = 1; ws.title = "W"
    ws.batch_get.return_value = [[[ln.get(h, "")] for ln in lines] for h in store._KEY_COLUMNS]
    return ws


def _store(lines=(), **kw):
    kw.setdefault("pacer", store._Pacer(0))
    sh = MagicMock()
    return sh, store.OpnameStore(sh, _ws(lines), **kw)


def _meta(tanggal="2026-10-25", toko="TOKO A", region="Sumatera"):
    return store.OpnameMeta(region, toko, tanggal, "Budi")


def _line(toko, periode, pid="P-1", bpom="NA1", region="Sumatera"):
    return {"Region": region, "Nama Toko (By Group)": toko, "Periode Stock Opname": periode, "PRODUCT ID": pid, "BPOM NUMBER": bpom}


def test_insert_writes_14_columns_region_toko_and_period_from_date():
    sh, opname = _store()
    row = v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5, additional=True)
    summary = opname.insert_rows(_meta(), [row])
    cells = sh.batch_update.call_args[0][0]["requests"][0]["appendCells"]["rows"][0]["values"]
    vals = {h: c["userEnteredValue"] for h, c in zip(store.HEADERS, cells)}
    assert len(cells) == len(store.HEADERS) == 14
    assert vals["Region"] == {"stringValue": "Sumatera"} and vals["Nama Toko (By Group)"] == {"stringValue": NAME}
    assert vals["Periode Stock Opname"] == {"stringValue": "2026-10"} and vals["Tanggal Stock Opname"] == {"stringValue": "2026-10-25"}
    assert vals["ISI (GR / ML)"] == {"stringValue": "10g"} and vals["QTY"] == {"numberValue": 5}
    assert (summary.inserted, summary.additional) == (1, 1)


def test_duplicates_are_per_region_toko_and_period():
    import pytest
    sh, opname = _store([_line("TOKO A", "2026-10")])
    row = v.ExcelRow(3, NAME, "p-1", "Produk", "10g", "na1", 5)
    with pytest.raises(store.DuplicateDataError) as exc:
        opname.insert_rows(_meta("2026-10-26"), [row])
    assert len(exc.value.conflicts) == 1 and not sh.batch_update.called
    assert opname.insert_rows(_meta("2026-10-26", "TOKO B"), [row]).inserted == 1   # toko lain
    assert opname.insert_rows(_meta("2026-11-02"), [row]).inserted == 1             # periode lain


def test_index_is_loaded_once_updated_by_own_writes_and_reloaded_after_ttl():
    import pytest
    clock = [0.0]
    sh, opname = _store(index_ttl=120, clock=lambda: clock[0])
    row = v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)
    ws = opname.ws
    opname.insert_rows(_meta(toko="A"), [row]); opname.insert_rows(_meta(toko="B"), [row])
    assert ws.batch_get.call_count == 1                                  # satu pembacaan untuk banyak submit
    assert len(ws.batch_get.call_args[0][0]) == 5                        # hanya 5 kolom kunci, bukan seluruh sheet
    with pytest.raises(store.DuplicateDataError):                        # penulisan sendiri langsung masuk indeks
        opname.insert_rows(_meta(toko="a"), [row])
    assert ws.batch_get.call_count == 1
    clock[0] = 121
    opname.insert_rows(_meta(toko="C"), [row])
    assert ws.batch_get.call_count == 2                                  # TTL lewat -> baca ulang


def test_new_spreadsheet_needs_no_initial_read():
    sh, opname = _store(assume_empty=True)
    opname.insert_rows(_meta(), [v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)])
    assert not opname.ws.batch_get.called


def _api_error(code):
    import gspread
    resp = MagicMock(); resp.status_code = code
    resp.json.return_value = {"error": {"code": code, "message": "x", "status": "X"}}
    return gspread.exceptions.APIError(resp)


def test_429_on_write_is_retried_with_backoff_but_5xx_is_not_and_forces_reread():
    import pytest
    sleeps = []
    sh, opname = _store(sleep=sleeps.append)
    sh.batch_update.side_effect = [_api_error(429), _api_error(429), {"ok": 1}]
    row = v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)
    assert opname.insert_rows(_meta(), [row]).inserted == 1
    assert sh.batch_update.call_count == 3 and sleeps == [2.0, 4.0]
    sh.batch_update.reset_mock(); sh.batch_update.side_effect = _api_error(503)       # hasil tidak pasti
    reads = opname.ws.batch_get.call_count
    with pytest.raises(Exception):
        opname.insert_rows(_meta(toko="Z"), [row])
    assert sh.batch_update.call_count == 1                                            # tidak diulang otomatis
    sh.batch_update.side_effect = None
    opname.insert_rows(_meta(toko="Y"), [row])
    assert opname.ws.batch_get.call_count == reads + 1                                # indeks dibaca ulang setelahnya


def test_write_gives_up_after_backoff_exhausted():
    import pytest
    sh, opname = _store(sleep=lambda s: None)
    sh.batch_update.side_effect = _api_error(429)
    with pytest.raises(Exception):
        opname.insert_rows(_meta(), [v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)])
    assert sh.batch_update.call_count == len(store._BACKOFF) + 1


def test_pacer_spaces_out_concurrent_writes_below_quota():
    now, slept = [0.0], []
    pacer = store._Pacer(1.1, clock=lambda: now[0], sleep=slept.append)
    for _ in range(5):
        pacer.wait()                                                       # lima penulisan di detik yang sama
    assert [round(x, 1) for x in slept] == [1.1, 2.2, 3.3, 4.4]            # tersebar ~54/menit (< kuota 60)


def test_concurrent_submits_of_same_toko_only_one_wins_and_other_tokos_all_write():
    import threading
    sh, opname = _store()
    row = v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)
    results = []

    def go(toko):
        try:
            results.append(opname.insert_rows(_meta(toko=toko), [row]).inserted)
        except store.DuplicateDataError:
            results.append("dup")
    threads = [threading.Thread(target=go, args=("SAMA",)) for _ in range(20)]
    threads += [threading.Thread(target=go, args=(f"S{i}",)) for i in range(40)]
    [t.start() for t in threads]; [t.join() for t in threads]
    assert results.count("dup") == 19 and results.count(1) == 41
    assert opname.ws.batch_get.call_count == 1
    written = [r for call in sh.batch_update.call_args_list for r in call[0][0]["requests"]]
    assert len(written) == 41                                     # setiap submit tepat satu permintaan, tidak hilang/ganda


def test_burst_of_100_submits_is_group_committed_into_few_requests():
    import threading, time
    sh, opname = _store(pacer=store._Pacer(0.2))                  # jeda antar-batch (di produksi 1.1 dtk)
    sh.batch_update.side_effect = lambda body: time.sleep(0.05)
    row = v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)
    done = []
    threads = [threading.Thread(target=lambda i=i: done.append(opname.insert_rows(_meta(toko=f"S{i}"), [row]).inserted))
               for i in range(100)]
    t0 = time.time(); [t.start() for t in threads]; [t.join() for t in threads]
    assert len(done) == 100 and sum(done) == 100
    calls = sh.batch_update.call_args_list
    assert sum(len(c[0][0]["requests"]) for c in calls) == 100
    assert len(calls) <= 10, len(calls)                           # bukan 100 request terpisah
    print(f"\n100 submit serentak -> {len(calls)} batchUpdate dalam {time.time() - t0:.1f}s")


def test_failure_of_a_group_commit_reaches_every_submit_in_it_and_forces_reread():
    import threading, pytest
    sh, opname = _store(pacer=store._Pacer(0.3))
    sh.batch_update.side_effect = _api_error(503)
    row = v.ExcelRow(3, NAME, "P-1", "Produk", "10g", "NA1", 5)
    outcomes = []

    def go(i):
        try:
            opname.insert_rows(_meta(toko=f"E{i}"), [row]); outcomes.append("ok")
        except Exception as exc:
            outcomes.append(type(exc).__name__)
    ts_ = [threading.Thread(target=go, args=(i,)) for i in range(6)]
    [t.start() for t in ts_]; [t.join() for t in ts_]
    assert outcomes == ["APIError"] * 6                           # tidak ada yang dilaporkan sukses palsu
    sh.batch_update.side_effect = None
    reads = opname.ws.batch_get.call_count
    assert opname.insert_rows(_meta(toko="E0"), [row]).inserted == 1     # tidak tercatat -> boleh submit ulang
    assert opname.ws.batch_get.call_count == reads + 1                       # setelah membaca ulang sheet


def test_existing_header_with_other_schema_is_never_overwritten():
    ws = MagicMock(); ws.row_values.return_value = ["header", "lain"]; ws.title = "W"
    sh = MagicMock(); sh.worksheet.return_value = ws
    import pytest
    with pytest.raises(store.SheetsError):
        store.ensure_worksheet(sh, "W")
    ws.update.assert_not_called()


def test_narrow_grid_is_widened_before_header_is_written():
    ws = MagicMock(); ws.row_values.return_value = []; ws.col_count = 10
    sh = MagicMock(); sh.worksheet.return_value = ws
    store.ensure_worksheet(sh, "W")
    ws.add_cols.assert_called_once_with(len(store.HEADERS) - 10)
    ws.update.assert_called_once()
    assert ws.update.call_args[0][0] == [store.HEADERS] and ws.update.call_args[0][1] == "A1:N1"


# ── SKU / BPOM tambahan + daftar SKU ────────────────────────────────────

import master_product as mp  # noqa: E402

PRODUCTS = [mp.Product("P-1", "Nama Master", "11ml"), mp.Product("Z-9", "Produk Baru", "25g"), mp.Product("NOPACK", "Tanpa Isi", "")]


def test_products_from_rows_dedupes_sorts_and_skips_incomplete():
    rows = [{"sku": "B-1", "product_name": "B", "pack_size": "1g"}, {"sku": "a-1", "product_name": "A", "pack_size": None},
            {"sku": "B-1", "product_name": "dup", "pack_size": "9g"}, {"sku": "C-1", "product_name": None, "pack_size": "1g"},
            {"sku": " ", "product_name": "X", "pack_size": "1g"}]
    products, skipped = mp.products_from_rows(rows)
    assert [p.label for p in products] == ["a-1 - A", "B-1 - B"] and skipped == 2 and products[1].pack_size == "1g"


def test_build_extra_uses_catalogue_name_and_isi_when_sku_is_official():
    extra, err = v.build_extra("p-1", " NA-X ", 3, PRODUCTS, SKUS)
    assert err is None and (extra.product_id, extra.product_name, extra.isi, extra.bpom_number, extra.qty) == ("P-1", "Produk", "10g", "NA-X", 3)
    extra, err = v.build_extra("Z-9", "NB-1", "0", PRODUCTS, SKUS)       # SKU di luar template -> dari master_product, QTY 0 sah
    assert err is None and (extra.product_name, extra.isi, extra.qty) == ("Produk Baru", "25g", 0)


def test_build_extra_rejections():
    for args, needle in ((("", "NA", 1), "SKU wajib"), (("Q-0", "NA", 1), "tidak ada di master"), (("Z-9", " ", 1), "BPOM wajib"),
                         (("Z-9", "NA", None), "QTY wajib"), (("Z-9", "NA", -1), "negatif"), (("Z-9", "NA", 1.5), "bulat"),
                         (("NOPACK", "NA", 1), "ISI"), (("Z-9", "N" * 201, 1), "maksimal")):
        extra, err = v.build_extra(*args, PRODUCTS, SKUS)
        assert extra is None and needle in err, (args, err)


def test_merge_extras_adds_rows_flags_additional_and_updates_missing():
    parsed = v.parse_excel(make([good()]), [NAME], SKUS)       # file hanya P-1/NA1 -> P-1/NA2 dan P-2/NB1 "missing"
    extras = [v.ExtraSku("P-1", "Produk", "10g", "NA2", 2),    # pasangan resmi yang hilang dari file -> mengisi
              v.ExtraSku("Z-9", "Produk Baru", "25g", "NB9", 4)]
    merged = v.merge_extras(parsed, extras, NAME, SKUS)
    assert merged.ok and [(r.product_id, r.bpom_number, r.additional, r.excel_row) for r in merged.rows[1:]] == [
        ("P-1", "NA2", False, 0), ("Z-9", "NB9", True, 0)]
    assert [k.bpom_number for _, k in merged.missing] == ["NB1"]
    assert v.merge_extras(parsed, [], NAME, SKUS) is parsed


def test_merge_extras_rejects_duplicates_all_or_nothing():
    parsed = v.parse_excel(make([good()]), [NAME], SKUS)
    dup_file = v.ExtraSku("p-1", "Produk", "10g", "na1", 1)
    dup_extra = [v.ExtraSku("Z-9", "Produk Baru", "25g", "NB9", 4), v.ExtraSku("z-9", "Produk Baru", "25g", "nb9", 5)]
    for extras, needle in (([dup_file], "baris 3 file Excel"), (dup_extra, "SKU tambahan #1")):
        merged = v.merge_extras(parsed, extras, NAME, SKUS)
        assert not merged.ok and needle in messages(merged)
    bad = v.parse_excel(make([good(qty=-1)]), [NAME], SKUS)
    assert v.merge_extras(bad, [dup_file], NAME, SKUS).issues == bad.issues     # file salah: tidak ditumpuk


def test_load_products_reads_sku_file_any_column_order_with_title_rows(tmp_path):
    import pytest
    path = _xlsx(tmp_path, "Daftar SKU.xlsx", [["EXPORT MASTER PRODUCT"], [None], ["Pack Size", "Product Name", "SKU", "Brand"],
                                               ["25ml", "Produk B", "B-1", "X"], ["5g", "Produk A", "a-1", "X"], ["1g", "dup", "B-1", "X"],
                                               [None, "Tanpa SKU", None, "X"], [None, None, None, None]])
    products, skipped = mp.load_products(path)
    assert [p.label for p in products] == ["a-1 - Produk A", "B-1 - Produk B"] and skipped == 1 and products[1].pack_size == "25ml"
    with pytest.raises(mp.SkuListMissing, match="belum tersedia"):
        mp.load_products(tmp_path / "tidak-ada.xlsx")
    with pytest.raises(mp.SkuListError, match="Header"):
        mp.load_products(_xlsx(tmp_path, "a.xlsx", [["Nama", "Harga"], ["x", 1]]))
    with pytest.raises(mp.SkuListError, match="tidak berisi SKU"):
        mp.load_products(_xlsx(tmp_path, "b.xlsx", [["SKU", "Product Name"]]))
    (only,) = mp.load_products(_xlsx(tmp_path, "c.xlsx", [["SKU", "Product Name"], ["Z-1", "Z"]]))[0]     # tanpa Pack Size boleh
    assert only.pack_size == ""


# ── daftar toko dengan filter kolom (Google Sheet: BA / Non BA) ─────────────

BA_HEAD = ["Region", "Nama Group Toko", "Nama Toko", "Kode", "Area", "PIC", "Cabang", "BA Non BA"]     # kolom H = indeks 7


def _ba(region, group, toko, flag):
    return [region, group, toko, "K", "A", "P", "C", flag]


def test_row_filter_keeps_only_exact_ba_rows_and_skips_validation_of_the_rest():
    rows = [BA_HEAD, _ba("Jawa", "ALFA", "Alfa 1", "BA"), _ba("Jawa", "ALFA", "Alfa 2", " ba "), _ba("Jawa", None, "Toko X", "Non BA"),
            _ba("Jawa", None, "Toko Y", ""), _ba("", None, "Rusak tapi Non BA", "Non BA"), _ba("Sumatera", None, "Toko Budi", "BA")]
    flt = r.RowFilter("BA", header_names=("BA Non BA",), column_letter="H")
    tokos = r.load_stores_from_values(rows, "sheet", row_filter=flt)
    assert [(t.region, t.name, t.outlets) for t in tokos] == [("Jawa", "ALFA", 2), ("Sumatera", "Toko Budi", 1)]   # "Non BA" bukan "BA"


def test_row_filter_falls_back_to_column_h_when_header_name_differs():
    head = [*BA_HEAD[:7], "Status BA"]
    rows = [head, _ba("Jawa", None, "A", "BA"), _ba("Jawa", None, "B", "Non BA")]
    flt = r.RowFilter("BA", header_names=("BA Non BA",), column_letter="H")
    assert [t.name for t in r.load_stores_from_values(rows, "sheet", row_filter=flt)] == ["A"]


def test_row_filter_errors_are_explicit():
    import pytest
    flt = r.RowFilter("BA", header_names=("BA Non BA",), column_letter="H")
    with pytest.raises(r.ReferenceDataError, match="Kolom filter tidak ditemukan"):
        r.load_stores_from_values([["Region", "Nama Toko"], ["Jawa", "A"]], "sheet", row_filter=flt)      # tidak ada kolom H
    with pytest.raises(r.ReferenceDataError, match="tidak berisi toko dengan BA"):
        r.load_stores_from_values([BA_HEAD, _ba("Jawa", None, "A", "Non BA")], "sheet", row_filter=flt)

    with pytest.raises(r.ReferenceDataError, match="Header 'sheet' tidak ditemukan.*ditemukan"):
        r.load_stores_from_values([["Wilayah", "Outlet"], ["Jawa", "A"]], "sheet", row_filter=flt)          # pesan memuat header yang ada


def test_no_match_error_names_the_filter_column_and_the_values_found():
    import pytest
    flt = r.RowFilter("BA", header_names=("BA Non BA",), column_letter="H")
    rows = [BA_HEAD, _ba("Jawa", None, "A", "Non BA"), _ba("Jawa", None, "B", "Non BA"), _ba("Jawa", None, "C", "Beauty Advisor"), _ba("Jawa", None, "D", "")]
    with pytest.raises(r.ReferenceDataError) as exc:
        r.load_stores_from_values(rows, "sheet", row_filter=flt)
    msg = str(exc.value)
    assert "kolom filter 'ba non ba' (kolom H)" in msg
    assert "'Non BA' (2)" in msg and "'Beauty Advisor' (1)" in msg and "'' (1)" in msg           # nilai yang ada, termasuk kosong
    assert "store_filter_value" in msg
    shifted = [["Region", "Nama Group Toko", "Nama Toko", "Kolom Baru", "x", "y", "z", "BA Non BA", "Status"], ["Jawa", None, "A", "", "", "", "", "Non BA", "BA"]]
    with pytest.raises(r.ReferenceDataError, match="kolom H"):                                      # kolom dicari lewat nama header, bukan posisi
        r.load_stores_from_values(shifted, "sheet", row_filter=flt)
    assert "TIDAK ditemukan" not in msg                                                             # header ada -> bukan cadangan kolom H

    renamed = [["Region", "Nama Group Toko", "Nama Toko", "Kode", "Area", "PIC", "Cabang", "SKP", "BA Flag"],
               ["Jawa", None, "A", "", "", "", "", "SKP", "BA"], ["Jawa", None, "B", "", "", "", "", "", "Non BA"]]
    with pytest.raises(r.ReferenceDataError) as exc2:                                               # kasus nyata: kolom H berganti makna
        r.load_stores_from_values(renamed, "sheet", row_filter=flt)
    m2 = str(exc2.value)
    assert "TIDAK ditemukan" in m2 and "'skp' (kolom H)" in m2 and "'SKP' (1)" in m2
    assert "H='skp'" in m2 and "I='ba flag'" in m2                                                  # semua header terlihat -> kolom pengganti langsung ketahuan


def test_index_letter_roundtrip():
    for letter in ("A", "H", "Z", "AA", "AZ", "BA"):
        assert r._index_letter(r._letter_index(letter)) == letter


def test_fetch_products_uses_a_parameterised_query_limited_to_the_given_brands():
    import pytest
    seen = {}

    class Job:
        def result(self):
            return [{"sku": "S-2", "product_name": "Dua", "pack_size": "5g"}, {"sku": "S-1", "product_name": "Satu", "pack_size": None},
                    {"sku": "S-2", "product_name": "dup", "pack_size": "9g"}, {"sku": "S-3", "product_name": None, "pack_size": "1g"}]

    class Client:
        def query(self, sql, job_config=None):
            seen["sql"], seen["params"] = sql, [(q.name, list(q.values)) for q in job_config.query_parameters]
            return Job()
    products, skipped = mp.fetch_products(Client(), "p.d.t", (" skintific ", "TIMEPHORIA", "FACERINNA", ""))
    assert [p.label for p in products] == ["S-1 - Satu", "S-2 - Dua"] and skipped == 1
    assert seen["params"] == [("brands", ["SKINTIFIC", "TIMEPHORIA", "FACERINNA"]),
                              ("exclude_status", ["DISCONTINUE", "DISCONTINUED", "NOT SOLD AT GT"])]   # nilai TIDAK disisipkan ke teks SQL
    assert "UNNEST(@brands)" in seen["sql"] and "UNNEST(@exclude_status)" in seen["sql"]
    assert "SKINTIFIC" not in seen["sql"] and "DISCONTINU" not in seen["sql"].upper() and "`p.d.t`" in seen["sql"]
    assert "status IS NULL OR" in seen["sql"]                                                      # status kosong tidak ikut dikecualikan
    mp.fetch_products(Client(), "p.d.t", ("A",), (" Obsolete ", ""))
    assert seen["params"][1] == ("exclude_status", ["OBSOLETE"])
    with pytest.raises(ValueError, match="tidak valid"):
        mp.fetch_products(Client(), "x; DROP TABLE y", ("A",))
    with pytest.raises(ValueError, match="brand master_product kosong"):
        mp.fetch_products(Client(), "p.d.t", ("", " "))


REAL_HEAD = ["store/group id", "kode toko", "nama toko", "region", "channel", "group", "budget q4", "skp", "tpm", "ba/ non ba",
             "program", "so", "store stock", "budget utilized"]                                       # header sheet BD per 2026-10-30


def _real(region, toko, group, flag, skp="SKP"):
    return ["ID1", "K1", toko, region, "MT", group, "0", skp, "", flag, "", "", "", ""]


def test_filter_matches_the_real_sheet_by_header_ignoring_spaces_and_slash_and_never_falls_back_to_a_position():
    import pytest
    flt = r.RowFilter("BA", header_names=("BA/ Non BA",))
    rows = [REAL_HEAD, _real("Jawa", "Alfa 1", "ALFAMART", "BA"), _real("Jawa", "Alfa 2", "ALFAMART", "BA"),
            _real("Jawa", "Toko X", None, "Non BA"), _real("Sumatera", "Toko Budi", None, "ba")]
    tokos = r.load_stores_from_values(rows, "sheet", row_filter=flt)
    assert [(t.region, t.name, t.outlets) for t in tokos] == [("Jawa", "ALFAMART", 2), ("Sumatera", "Toko Budi", 1)]
    for alias in ("BA Non BA", "ba/non ba", "BA / NON BA"):
        assert len(r.load_stores_from_values(rows, "sheet", row_filter=r.RowFilter("BA", header_names=(alias,)))) == 2
    no_marker = [[h for h in REAL_HEAD if h != "ba/ non ba"], *[[v for i, v in enumerate(row) if i != 9] for row in rows[1:]]]
    with pytest.raises(r.ReferenceDataError, match="Kolom filter tidak ditemukan"):                   # tidak ada kolom H/J diam-diam
        r.load_stores_from_values(no_marker, "sheet", row_filter=flt)
