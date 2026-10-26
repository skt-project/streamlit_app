"""Tes skrip persiapan data BD (tools/prepare_reference.py) memakai file buatan di folder sementara."""
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import prepare_reference as pr  # noqa: E402
import reference_data as r  # noqa: E402

CURRENT = r.load_catalogue(ROOT / "template" / "Contoh Excel Response.xlsx")


def xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


@pytest.fixture
def sandbox(tmp_path):
    """Salinan struktur proyek (reference + template) supaya --install tidak menyentuh file asli."""
    (tmp_path / "reference").mkdir(); (tmp_path / "template").mkdir()
    shutil.copyfile(ROOT / "template" / "Contoh Excel Response.xlsx", tmp_path / "template" / "Contoh Excel Response.xlsx")
    return tmp_path


def run(sandbox, *argv):
    return pr.main(["--root", str(sandbox), *map(str, argv)])


HEAD = ["Region", "Nama Group Toko", "Nama Toko"]


def test_store_file_with_title_rows_blank_tail_and_extra_columns_is_accepted(tmp_path):
    path = xlsx(tmp_path / "s.xlsx", [["DAFTAR TOKO BD"], [], ["Area", "Nama Toko", "Region", "Nama Group Toko", "Catatan"],
                                      ["x", "Alfa 1", "Jawa", "ALFAMART", "-"], ["x", "Alfa 2", "Jawa", "ALFAMART", None],
                                      ["x", "Toko Budi", "Sumatera", None, None], [None] * 5])
    rep, tokos = pr.check_stores(path)
    assert rep.ok and [(t.region, t.name, t.outlets) for t in tokos] == [("Jawa", "ALFAMART", 2), ("Sumatera", "Toko Budi", 1)]
    assert any("2 pilihan toko valid di 2 region (1 grup menggabungkan 2 baris toko; 1 toko tanpa grup)" in i for i in rep.info)


def test_store_problems_are_all_listed(tmp_path):
    path = xlsx(tmp_path / "s.xlsx", [HEAD, ["", None, "A"], ["Jawa", None, ""], ["Jawa", "ALFA", "x"], ["Jawa", None, "ALFA"]])
    rep, _ = pr.check_stores(path)
    text = "\n".join(rep.errors)
    assert not rep.ok and "baris 2" in text and "baris 3" in text and "grup dan sebagai toko tunggal" in text


def test_missing_required_column_is_an_error_and_extra_sheets_are_warned(tmp_path):
    rep, _ = pr.check_stores(xlsx(tmp_path / "a.xlsx", [["Nama Toko"], ["x"]]))
    assert not rep.ok and any("Region" in e for e in rep.errors)
    wb = openpyxl.Workbook(); wb.active.append(HEAD); wb.active.append(["Jawa", None, "A"]); wb.create_sheet("Lain")
    wb.save(tmp_path / "b.xlsx")
    rep, _ = pr.check_stores(tmp_path / "b.xlsx")
    assert rep.ok and any("HANYA sheet pertama" in w for w in rep.warnings)


def test_ungrouped_duplicate_names_in_one_region_are_warned_and_shared_names_noted(tmp_path):
    path = xlsx(tmp_path / "s.xlsx", [HEAD, ["Jawa", None, "Toko A"], ["Jawa", None, "toko a"], ["Sumatera", None, "Toko A"]])
    rep, tokos = pr.check_stores(path)
    assert rep.ok and len(tokos) == 2 and any("lebih dari sekali di region yang sama" in w for w in rep.warnings)
    assert any("lebih dari satu region" in i for i in rep.info)


def test_install_only_when_everything_passes(sandbox, tmp_path):
    bad = xlsx(tmp_path / "bad.xlsx", [HEAD, ["", None, "A"]])
    assert run(sandbox, "--stores", bad, "--install") == 1
    assert not (sandbox / "reference" / "Daftar Store.xlsx").exists()
    good = xlsx(tmp_path / "good.xlsx", [HEAD, ["Jawa", None, "A"]])
    assert run(sandbox, "--stores", good) == 0 and not (sandbox / "reference" / "Daftar Store.xlsx").exists()   # tanpa --install
    assert run(sandbox, "--stores", good, "--install") == 0
    assert r.load_stores(sandbox / "reference" / "Daftar Store.xlsx")[0].name == "A"


def test_catalogue_in_template_format_replaces_sheet1_and_keeps_header_and_other_sheets(sandbox, tmp_path):
    first = CURRENT[0]
    path = xlsx(tmp_path / "c.xlsx", [
        ["PRODUCT ID", "PRODUCT NAME", "ISI (GR / ML)", "BPOM NUMBER"],
        [first.product_id, first.product_name, first.isi, first.bpom_number],
        ["ZZ-1", "PRODUK BARU", "10ml", "NA11260100001"], [None] * 4])
    assert run(sandbox, "--catalogue", path, "--install") == 0
    tpl = sandbox / "template" / "Contoh Excel Response.xlsx"
    new = r.load_catalogue(tpl)
    assert [(s.product_id, s.bpom_number) for s in new] == [(first.product_id, first.bpom_number), ("ZZ-1", "NA11260100001")]
    wb = openpyxl.load_workbook(tpl)
    assert wb.sheetnames == ["Sheet1", "Distributor Name", "SKU", "Sheet3"]            # sheet lain tidak disentuh
    header = [c.value for c in wb["Sheet1"][2]]
    assert header[1:] == list(pr.val.EXPECTED_HEADERS[1:]) and header[0] in (pr.val.EXPECTED_HEADERS[0], pr.val.LEGACY_FIRST_HEADER)
    assert wb["Sheet1"]["F1"].value == "Filled by Distributor"


def test_catalogue_in_bd_format_is_completed_from_current_catalogue_and_master(tmp_path):
    first = CURRENT[0]
    path = xlsx(tmp_path / "c.xlsx", [
        ["SKU", "BPOM Product Name", "BPOM Number", "Expired Date"],
        [first.product_id, "nama BPOM bebas", "NA11299999999", None],                  # SKU dikenal -> nama + ISI dari katalog
        ["ZZ-9", "x", "NA11200000009", None],                                          # SKU baru -> dari master
        ["QQ-0", "x", "NA11200000010", None]])                                         # tidak dikenal di mana pun
    import master_product as mp
    rep, skus = pr.read_catalogue_input(path, CURRENT, [mp.Product("ZZ-9", "PRODUK ZZ", "5g")])
    assert not rep.ok and any("QQ-0" in e for e in rep.errors)
    assert [(s.product_id, s.product_name, s.isi) for s in skus] == [
        (first.product_id, first.product_name, first.isi), ("ZZ-9", "PRODUK ZZ", "5g")]
    rep2, _ = pr.read_catalogue_input(path, CURRENT, None)
    assert any("--bq" in e for e in rep2.errors)


def test_catalogue_rejects_duplicates_missing_bpom_and_warns_on_odd_bpom_and_removals(tmp_path):
    a = CURRENT[0]
    path = xlsx(tmp_path / "c.xlsx", [
        ["PRODUCT ID", "PRODUCT NAME", "ISI (GR / ML)", "BPOM NUMBER"],
        [a.product_id, a.product_name, a.isi, a.bpom_number], [a.product_id, a.product_name, a.isi, a.bpom_number.lower()],
        [a.product_id, a.product_name, a.isi, None], [a.product_id, a.product_name, a.isi, "BPOM-SALAH"]])
    rep, _ = pr.read_catalogue_input(path, CURRENT)
    assert any("ganda" in e for e in rep.errors) and any("wajib diisi" in e for e in rep.errors)
    assert any("tidak berpola" in w for w in rep.warnings) and any("Hilang dari katalog baru" in w for w in rep.warnings)


def test_real_current_catalogue_roundtrips_without_changes(sandbox):
    """Katalog saat ini dimasukkan lagi -> tidak ada perubahan (+0/-0)."""
    rows = [["PRODUCT ID", "PRODUCT NAME", "ISI (GR / ML)", "BPOM NUMBER"]] + [[s.product_id, s.product_name, s.isi, s.bpom_number] for s in CURRENT]
    path = xlsx(sandbox / "same.xlsx", rows)
    rep, skus = pr.read_catalogue_input(path, CURRENT)
    assert rep.ok and len(skus) == len(CURRENT) and any("+0 baru, -0 hilang, 0 nama/ISI berubah" in i for i in rep.info)
