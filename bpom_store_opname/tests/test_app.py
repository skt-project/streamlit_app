"""Tes UI end-to-end (Streamlit AppTest) dengan Google Sheets di-mock. Tidak menyentuh spreadsheet asli.

Daftar toko di sini adalah DATA BUATAN di folder sementara; tidak ada yang ditulis ke reference/.
"""
import sys
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import app  # noqa: E402
import master_product as mp  # noqa: E402
import reference_data as r  # noqa: E402
import sheets_store as store  # noqa: E402
import validation as v  # noqa: E402
from template_builder import build_template  # noqa: E402

TEMPLATE = ROOT / "template" / "Contoh Excel Response.xlsx"
REGION = "SUMATERA"
TOKO = "GRUP ALFA"          # nama grup (2 cabang di daftar BD) -> pilihan pertama di dropdown (urut abjad)
TOKO2 = "TOKO MAJU"
SKUS = r.load_catalogue(TEMPLATE)
# master_product buatan: satu SKU yang ada di template (nama beda), satu SKU baru di luar template
PRODUCTS = [mp.Product(SKUS[0].product_id, "NAMA MASTER BEDA", "999ml"), mp.Product("ZZ-NEW-01", "PRODUK BARU ZZ", "25ml")]


def _run():
    import app as module
    module.main()


def _store_file(tmp_path, rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nama Toko", "Nama Group Toko", "Region"])
    for row in rows or [["TOKO MAJU", None, REGION], ["Cabang 1", TOKO, REGION], ["Cabang 2", TOKO, REGION]]:
        ws.append(row)
    path = tmp_path / "Daftar Store.xlsx"
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def wiring(monkeypatch, tmp_path):
    """Credential palsu, Google Sheets palsu, daftar store buatan."""
    monkeypatch.setattr(app, "load_config", lambda: app.Config(
        sheet_name="T", worksheet_title="T", sheet_id="FAKE_ID", share_with=(), template_path=TEMPLATE))
    monkeypatch.setattr(app, "get_gspread_client", lambda: MagicMock())
    monkeypatch.setattr(app, "load_products", lambda: (PRODUCTS, 0, None))
    monkeypatch.setattr(app, "STORE_FILE", _store_file(tmp_path))
    monkeypatch.setattr(store, "WRITE_INTERVAL", 0.0)       # tanpa jeda penulisan di tes
    for cached in (app.load_tokos_by_region, app.load_skus, app.get_opname_store):
        cached.clear()
    yield
    for cached in (app.load_tokos_by_region, app.get_opname_store):
        cached.clear()


def _columns(lines):
    """Isi balasan ws.batch_get (4 kolom kunci) dari daftar baris {header: nilai}."""
    return [[[ln.get(h, "")] for ln in lines] for h in store._KEY_COLUMNS]


def _fake_sheet(monkeypatch, existing=None):
    sh, ws = MagicMock(), MagicMock()
    ws.id, ws.title = 7, "W"
    ws.batch_get.return_value = _columns(existing or [])
    sh.ws = ws
    monkeypatch.setattr(store, "open_spreadsheet", lambda *a, **k: (sh, False))
    monkeypatch.setattr(store, "ensure_worksheet", lambda *a, **k: ws)
    return sh


def _filled(region, toko, qty=lambda i: i):
    wb = openpyxl.load_workbook(BytesIO(build_template(TEMPLATE.read_bytes(), region, toko, SKUS)))
    ws = wb["Sheet1"]
    for i in range(3, ws.max_row + 1):
        ws.cell(i, 6).value = qty(i)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fill_form(at, upload=None):
    at.selectbox[0].select(REGION)
    at.run()
    at.selectbox[1].select_index(0)
    at.date_input[0].set_value(__import__("datetime").date(2026, 10, 25))
    at.text_input[0].input("Budi")
    if upload is not None:
        at.file_uploader[0].upload("isi.xlsx", upload, "application/octet-stream")
    at.run()
    return at


def _click(at, label):
    next(b for b in at.button if b.label == label).click()
    at.run()


def _add_extra(at, product_index, bpom, qty):
    at.selectbox[2].select_index(product_index)
    next(t for t in at.text_input if t.label == "Nomor BPOM").input(bpom)
    at.number_input[0].set_value(qty)
    _click(at, "Tambah ke daftar")


def _at():
    at = AppTest.from_function(_run, default_timeout=60)
    at.run()
    return at


def test_missing_store_list_shows_message_and_no_form(monkeypatch, tmp_path):
    monkeypatch.setattr(app, "STORE_FILE", tmp_path / "tidak-ada.xlsx")
    app.load_tokos_by_region.clear()
    at = _at()
    assert not at.exception
    assert any("Daftar toko belum tersedia" in w.value for w in at.warning)
    assert len(at.selectbox) == 0 and len(at.button) == 0


def test_two_step_selection_store_label_and_download_button():
    at = _at()
    assert not at.exception
    assert [s.label for s in at.selectbox] == ["Region", "Toko", "SKU"] and at.selectbox[1].disabled
    assert not at.download_button            # belum ada store -> belum ada template
    at.selectbox[0].select(REGION)
    at.run()
    store_box = at.selectbox[1]
    assert not store_box.disabled and len(store_box.options) == 2
    assert store_box.options[0] == "GRUP ALFA  ·  grup (2 toko)"        # grup = SATU pilihan atas nama grup
    assert store_box.options[1] == "TOKO MAJU"                           # tanpa grup -> nama toko
    at.selectbox[1].select_index(0)
    at.run()
    assert len(at.download_button) == 1


def test_submit_valid_file_writes_one_atomic_batch(monkeypatch):
    sh = _fake_sheet(monkeypatch)
    at = _fill_form(_at(), _filled(REGION, TOKO))
    assert not at.exception and not at.error
    assert any("File valid" in s.value for s in at.success)
    _click(at, "Submit")
    assert not at.exception and not at.error
    assert sh.batch_update.call_count == 1                           # SATU batchUpdate, atomik
    rows = sh.batch_update.call_args[0][0]["requests"][0]["appendCells"]["rows"]
    assert len(rows) == len(SKUS)
    cells = {h: c["userEnteredValue"] for h, c in zip(store.HEADERS, rows[0]["values"])}
    assert cells["Region"] == {"stringValue": REGION} and cells["Nama Toko (By Group)"] == {"stringValue": TOKO}
    assert len(rows[0]["values"]) == 14
    assert cells["Periode Stock Opname"] == {"stringValue": "2026-10"} and cells["created_by"] == {"stringValue": "Budi"}
    assert isinstance(cells["QTY"].get("numberValue"), int)
    assert any("Berhasil menyimpan" in s.value for s in at.success)


def test_invalid_row_blocks_everything_with_row_number_and_nothing_written(monkeypatch):
    sh = _fake_sheet(monkeypatch)
    data = _filled(REGION, TOKO, qty=lambda i: -4 if i == 5 else 1)
    at = _fill_form(_at(), data)
    assert any("Baris 5" in m.value and "negatif" in m.value for m in at.markdown)
    _click(at, "Submit")
    assert not at.exception and sh.batch_update.call_count == 0
    assert any("belum valid" in e.value for e in at.error)


def test_file_of_another_toko_is_rejected(monkeypatch):
    sh = _fake_sheet(monkeypatch)
    at = _fill_form(_at(), _filled(REGION, TOKO2))        # file milik toko lain, bukan toko yang dipilih
    assert any("bukan toko yang dipilih" in m.value for m in at.markdown)
    _click(at, "Submit")
    assert sh.batch_update.call_count == 0


def test_duplicate_in_sheet_rejects_whole_input(monkeypatch):
    first = SKUS[0]
    line = {"Region": REGION, "Nama Toko (By Group)": TOKO, "Periode Stock Opname": "2026-10", "PRODUCT ID": first.product_id,
            "BPOM NUMBER": first.bpom_number}
    sh = _fake_sheet(monkeypatch, existing=[line])
    at = _fill_form(_at(), _filled(REGION, TOKO))
    _click(at, "Submit")
    assert sh.batch_update.call_count == 0
    assert any("hubungi BD" in e.value for e in at.error)


def test_form_fields_are_all_required(monkeypatch):
    sh = _fake_sheet(monkeypatch)
    at = _at()
    _click(at, "Submit")
    msgs = " ".join(e.value for e in at.error)
    for needle in ("Region wajib", "Toko wajib", "Tanggal", "Nama PIC", "File Excel"):
        assert needle in msgs
    assert sh.batch_update.call_count == 0


def test_large_company_uses_search_box_and_stays_fast(tmp_path, monkeypatch):
    import time
    rows = [[f"TOKO {i:06d}", None, REGION] for i in range(20000)]
    monkeypatch.setattr(app, "STORE_FILE", _store_file(tmp_path, rows))
    app.load_tokos_by_region.clear()
    t0 = time.perf_counter()
    at = _at()
    cold = time.perf_counter() - t0
    at.selectbox[0].select(REGION)
    t1 = time.perf_counter()
    at.run()
    warm = time.perf_counter() - t1
    print(f"\n20000 store: cold load {cold:.2f}s, rerun setelah pilih company {warm:.2f}s")
    assert cold < 15 and warm < 5
    search = next(t for t in at.text_input if t.label == "Cari Toko")
    assert at.selectbox[1].options == [] and at.selectbox[1].disabled         # belum mengetik -> tidak ada opsi dikirim
    search.input("to").run()
    assert at.selectbox[1].options == []                                      # < 3 huruf
    next(t for t in at.text_input if t.label == "Cari Toko").input("toko 01999").run()
    opts = list(at.selectbox[1].options)
    assert 0 < len(opts) <= app.STORE_SEARCH_RESULTS and all("01999" in o for o in opts)
    assert any("019990" in o for o in opts)
    at.selectbox[1].select_index(0).run()
    assert len(at.download_button) == 1


def test_company_below_threshold_keeps_plain_dropdown(tmp_path, monkeypatch):
    rows = [[f"TOKO {i:05d}", None, REGION] for i in range(app.STORE_SEARCH_THRESHOLD)]
    monkeypatch.setattr(app, "STORE_FILE", _store_file(tmp_path, rows))
    app.load_tokos_by_region.clear()
    at = _at()
    at.selectbox[0].select(REGION).run()
    assert not any(t.label == "Cari Toko" for t in at.text_input)
    assert len(at.selectbox[1].options) == app.STORE_SEARCH_THRESHOLD


def test_sheets_calls_per_submit_are_minimal_and_handle_is_reused(monkeypatch):
    """Submit kedua tidak membuka spreadsheet/worksheet lagi dan tidak membaca sheet lagi (indeks di memori)."""
    sh = _fake_sheet(monkeypatch)
    opened = []
    real = store.open_spreadsheet
    monkeypatch.setattr(store, "open_spreadsheet", lambda *a, **k: (opened.append(1), real(*a, **k))[1])
    for _ in range(2):
        at = _fill_form(_at(), _filled(REGION, TOKO))
        _click(at, "Submit")
    # submit kedua: toko yang sama + periode sama -> ditolak dari INDEKS, tanpa membaca sheet lagi
    assert any("hubungi BD" in e.value for e in at.error)
    assert len(opened) == 1 and sh.ws.batch_get.call_count == 1 and sh.batch_update.call_count == 1



# ── fitur tambah SKU / BPOM langsung (SKU dari master_product) ──────────────

def test_extra_form_lists_sku_dash_product_name_and_adds_to_list():
    at = _at()
    assert list(at.selectbox[2].options) == [f"{SKUS[0].product_id} - NAMA MASTER BEDA", "ZZ-NEW-01 - PRODUK BARU ZZ"]
    _add_extra(at, 1, " NA999 ", 4)
    assert not at.exception and not at.error
    assert any("ZZ-NEW-01" in m.value and "NA999" in m.value and "ISI 25ml" in m.value for m in at.markdown)
    assert any("Ditambahkan" in x.value for x in at.success)


def test_extra_form_validation_errors():
    at = _at()
    _click(at, "Tambah ke daftar")
    assert any("SKU wajib" in e.value for e in at.error)
    at.selectbox[2].select_index(1); _click(at, "Tambah ke daftar")
    assert any("Nomor BPOM wajib" in e.value for e in at.error)


def test_extras_are_saved_with_file_rows_in_one_batch(monkeypatch):
    sh = _fake_sheet(monkeypatch)
    at = _at()
    _add_extra(at, 1, "NA999", 4)        # SKU baru di luar template
    _add_extra(at, 0, "NA-BARU-77", 0)   # SKU template + BPOM baru, QTY 0 sah
    at = _fill_form(at, _filled(REGION, TOKO))
    assert not at.error and any("2 SKU/BPOM tambahan" in x.value for x in at.success)
    _click(at, "Submit")
    assert not at.exception and not at.error and sh.batch_update.call_count == 1
    rows = sh.batch_update.call_args[0][0]["requests"][0]["appendCells"]["rows"]
    assert len(rows) == len(SKUS) + 2
    last = [{h: c["userEnteredValue"] for h, c in zip(store.HEADERS, r["values"])} for r in rows[-2:]]
    assert last[0]["PRODUCT ID"] == {"stringValue": "ZZ-NEW-01"} and last[0]["PRODUCT NAME"] == {"stringValue": "PRODUK BARU ZZ"}
    assert last[0]["ISI (GR / ML)"] == {"stringValue": "25ml"} and last[0]["QTY"] == {"numberValue": 4}
    assert last[0]["Region"] == {"stringValue": REGION} and last[0]["Nama Toko (By Group)"] == {"stringValue": TOKO}
    # SKU yang ada di template memakai nama + ISI resmi template, bukan master_product
    assert last[1]["PRODUCT NAME"] == {"stringValue": SKUS[0].product_name}
    assert last[1]["ISI (GR / ML)"] == {"stringValue": SKUS[0].isi} and last[1]["QTY"] == {"numberValue": 0}
    assert not at.session_state["extras"]            # daftar dikosongkan setelah sukses


def test_extra_duplicating_file_row_blocks_submit(monkeypatch):
    sh = _fake_sheet(monkeypatch)
    at = _at()
    _add_extra(at, 0, SKUS[0].bpom_number, 3)        # pasangan resmi yang sudah ada di file
    at = _fill_form(at, _filled(REGION, TOKO))
    assert any("sudah ada" in m.value and "SKU tambahan #1" in m.value for m in at.markdown)
    _click(at, "Submit")
    assert sh.batch_update.call_count == 0


def test_extra_can_be_removed():
    at = _at()
    _add_extra(at, 1, "NA999", 4)
    _click(at, "Hapus")
    assert not at.session_state["extras"]


def test_missing_sku_list_disables_only_the_extra_feature(monkeypatch, tmp_path):
    monkeypatch.undo()                                                    # pakai load_products asli, bukan tiruan
    monkeypatch.setattr(app, "load_config", lambda: app.Config(
        sheet_name="T", worksheet_title="T", sheet_id="FAKE_ID", share_with=(), template_path=TEMPLATE))
    monkeypatch.setattr(app, "STORE_FILE", _store_file(tmp_path))
    monkeypatch.setattr(app, "SKU_LIST_FILE", tmp_path / "tidak-ada.xlsx")
    app.load_tokos_by_region.clear(); app._read_products.clear()
    at = _at()
    assert not at.exception
    assert any("Daftar SKU belum tersedia" in w.value for w in at.warning)
    assert len(at.selectbox) == 2 and len(at.download_button) == 0     # form region/toko tetap jalan


def test_app_has_no_bigquery_dependency():
    source = (ROOT / "app.py").read_text(encoding="utf-8")
    assert "import bigquery" not in source and "google.cloud" not in source
    assert "google-cloud-bigquery" not in (ROOT / "requirements.txt").read_text(encoding="utf-8")


def test_sku_list_file_feeds_the_dropdown(monkeypatch, tmp_path):
    monkeypatch.undo()
    monkeypatch.setattr(app, "load_config", lambda: app.Config(
        sheet_name="T", worksheet_title="T", sheet_id="FAKE_ID", share_with=(), template_path=TEMPLATE))
    monkeypatch.setattr(app, "STORE_FILE", _store_file(tmp_path))
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Pack Size", "SKU", "Product Name"]); ws.append(["25ml", "ZZ-NEW-01", "PRODUK BARU ZZ"]); ws.append(["5g", "AA-1", "PRODUK A"])
    wb.save(tmp_path / "Daftar SKU.xlsx")
    monkeypatch.setattr(app, "SKU_LIST_FILE", tmp_path / "Daftar SKU.xlsx")
    app.load_tokos_by_region.clear(); app._read_products.clear()
    at = _at()
    assert not at.exception and list(at.selectbox[2].options) == ["AA-1 - PRODUK A", "ZZ-NEW-01 - PRODUK BARU ZZ"]
